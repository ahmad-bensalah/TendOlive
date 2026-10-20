"""Load authoritative PDF CVs and indexed structured JSON declarations.

CVs come only from data/cvs. Project, stack and client records are filtered by
request-local Company/Demo mode. Draft and invalid records are excluded.
Text-reading helpers remain available for CV parsing; uploads are not an ingestion path.
"""
import re
import unicodedata
from dataclasses import dataclass, field
from pathlib import Path

from . import config
from . import cvparse
from .context import current
from . import records
import json

FOLDER_TYPES = {
    "authoritative_cvs": "cv",
    "projects": "past_project",
    "tech_stacks": "tech_stack",
    "clients": "client_portfolio",
    "tools": "tool",
}
DOC_TYPES = set(FOLDER_TYPES.values())
EXTENSIONS = {".md", ".txt", ".pdf", ".docx"}
MIN_WORDS = 20  # files with less text are treated as empty or unreadable


def active_folders() -> list[str]:
    """Folders to load: all of them, or only those listed in RAG_FOLDERS."""
    if not current().folders:
        return list(FOLDER_TYPES)
    return [f for f in FOLDER_TYPES if f in current().folders]


def source_files(data_dir: Path):
    """CVs always come from CV_DIR; other documents come from data/."""
    for folder in active_folders():
        root = config.CV_DIR if folder == "authoritative_cvs" else data_dir / folder
        if folder == "authoritative_cvs" and not root.is_dir():
            raise FileNotFoundError(f"Authoritative CV folder not found: {root}")
        for path in sorted(root.glob("*")):
            if not path.is_file() or path.stem.upper() == "README":
                continue
            if folder == "authoritative_cvs":
                if path.suffix.lower() != ".pdf":
                    continue
            elif path.suffix.lower() != ".json":
                continue
            yield folder, path


@dataclass
class Document:
    doc_id: str
    doc_type: str
    title: str
    summary: str
    text: str
    path: str = ""
    record: dict = field(default_factory=dict)


def data_fingerprint(data_dir: Path | None = None) -> str:
    """Content hash of active files, including the authoritative CV source."""
    import hashlib
    data_dir = Path(data_dir or current().data_dir)
    h = hashlib.sha1()
    h.update(str(config.CV_DIR.resolve()).encode())
    for folder, path in source_files(data_dir):
        h.update(f"{folder}/{path.name}\n".encode())
        h.update(path.read_bytes())
    return h.hexdigest()[:16]


def read_file(path: Path) -> str:
    """Plain text of a supported file (raises on unreadable files)."""
    return _read(path).strip()


_ACCENTS = {"´": "\u0301", "`": "\u0300", "ˆ": "\u0302", "¨": "\u0308"}
# LaTeX "T1" font codes that a PDF without a Unicode table leaves in the text:
# ligatures ("Air\x1dow" is "Airflow"), dashes ("2020 \x15 2023") and quotes.
_T1_CODES = str.maketrans({"\x10": "“", "\x11": "”", "\x12": "„", "\x13": "«", "\x14": "»",
                           "\x15": "–", "\x16": "—", "\x17": "", "\x1b": "ff", "\x1c": "fi",
                           "\x1d": "fl", "\x1e": "ffi", "\x1f": "ffl"})


_KNOWN = set()


def _known_words() -> set:
    """One-word names from the vocabulary and the CV section titles (for kerning repairs)."""
    if not _KNOWN:
        from . import taxonomy as tx
        names = [a for table in (tx.SKILLS, tx.CERTIFICATIONS, tx.ROLES) for aliases in table.values()
                 for a in aliases] + [n for names in cvparse.SECTION_NAMES.values() for n in names]
        _KNOWN.update(w for n in names for w in n.split() if w.isalpha() and len(w) > 1)
        _KNOWN.update({"awards", "associate", "architect", "analyst", "administrator", "analytics",
                       "automation", "academic", "advanced", "application", "applications"})
    return _KNOWN


def clean_pdf_text(text: str) -> str:
    """Repair text that PDF export broke (common in CVs made with LaTeX).

    "Facult´ e" -> "Faculté", "¿15TB" -> ">15TB", "¡200ms" -> "<200ms", "3 Ö" -> "3×",
    "Air\x1dow" -> "Airflow", "2020 \x15 2023" -> "2020 – 2023", "T eam" -> "Team",
    bullet and arrow symbols, a literal "textbf" or "ill" (a broken \\hfill before dates),
    Markdown "**", the "L ATEX" logo, math left as text ("Kafka $ rightarrow$ Spark" ->
    "Kafka → Spark", "3$ times$" -> "3×") and a "%" left before the first skills line.
    """
    text = text.translate(_T1_CODES)
    # LaTeX math printed as its source, often cut at a line end: "Kafka $\nrightarrow$ Spark".
    text = re.sub(r"\s*\$\s*\\?rightarrow\s*\$\s*", " → ", text)
    text = re.sub(r"(\d)\s*\$\s*\\?times\s*\$\s*", r"\1× ", text)
    # "% Cloud Platforms: AWS, ..." (a "%" that is not a percentage, at the start of a "Label:" line).
    text = re.sub(r"(?m)^%[ \t]+(?=[^\W\d_][^:\n]{0,60}:)", "", text)
    text = text.replace("\x88", "").replace("\x19", " -> ").replace("**", "")
    text = re.sub(r"\s+ill\s+(?=(?:[^\W\d_]{3,9}\.?\s*)?(?:19|20)\d\d)", " ", text)
    # Kerning read as a space after a capital letter: "T eam", "W eb", "P ython".
    text = re.sub(r"(?<![\w-])([FPTVWY]) (?=[a-z]{2})", r"\1", text)
    # After other capitals ("A wards", "A WS") only when the joined word is known,
    # so "A web app" stays as it is.
    text = re.sub(r"(?<![\w-])([A-Z]) ([A-Za-z]{2,})\b",
                  lambda m: m.group(1) + m.group(2) if (m.group(1) + m.group(2)).lower() in _known_words()
                  else m.group(0), text)
    text = re.sub(r"\\?textbf\{?", "", text)
    text = text.replace("L ATEX", "LaTeX")
    text = re.sub(r"\s*¿\s*(?=[\d$€£])", " >", text)
    text = re.sub(r"\s*¡\s*(?=[\d$€£])", " <", text)
    text = re.sub(r"(\d)\s?Ö\s?", r"\1× ", text)
    text = re.sub(r"([´`ˆ¨])\s?([A-Za-z])",
                  lambda m: unicodedata.normalize("NFC", m.group(2) + _ACCENTS[m.group(1)]), text)
    return text


def _read(path: Path) -> str:
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        from pypdf import PdfReader
        text = clean_pdf_text("\n".join((p.extract_text() or "") for p in PdfReader(str(path)).pages))
        # One sentence or bullet item per line, so a quote is never the tail of a sentence.
        return cvparse.join_wrapped_lines(text)
    if suffix == ".docx":
        from docx import Document as DocxFile  # package: python-docx
        doc = DocxFile(str(path))
        parts = [p.text for p in doc.paragraphs]
        for table in doc.tables:  # many CV templates put skills in tables
            for row in table.rows:
                parts.append(" | ".join(cell.text.strip() for cell in row.cells))
        return "\n".join(parts)
    return path.read_text(encoding="utf-8", errors="ignore")


def _parse_header(text: str) -> dict:
    """Read the optional 'Title:', 'Type:', 'Summary:' lines at the top of a file."""
    meta = {}
    for line in text.splitlines()[:5]:
        for key in ("Title", "Type", "Summary"):
            if line.startswith(f"{key}:"):
                meta[key.lower()] = line.split(":", 1)[1].strip()
    return meta


def load_documents(data_dir: Path | None = None) -> list[Document]:
    data_dir = Path(data_dir or current().data_dir)
    docs, seen = [], set()
    for folder, path in source_files(data_dir):
        if folder != "authoritative_cvs":
            record = json.loads(path.read_text(encoding="utf-8"))
            origin = "demo" if current().mode == "demo" else "user_entered"
            if record.get("origin") != origin or record.get("status") != "indexed":
                continue
            if not records.validate_record(record, records.all_records(data_dir))["valid"]:
                continue
            text = "\n".join(records.statements(record))
            docs.append(Document(record["id"], record["type"], record["title"], record["description"],
                                 text, path.relative_to(data_dir).as_posix(), record))
            continue
        default_type = FOLDER_TYPES[folder]
        # Keep API document IDs stable while preserving source filenames on disk.
        doc_id = "CV_" + path.name if folder == "authoritative_cvs" else path.name
        if path.suffix.lower() not in EXTENSIONS or path.stem.upper() == "README":
            continue
        if doc_id in seen:
            print(f"[loader] skipped (same file name already loaded): {folder}/{path.name}")
            continue
        try:
            text = _read(path).strip()
        except Exception as e:  # a broken file must not stop indexing
            print(f"[loader] skipped (cannot read): {path.name}: {e}")
            continue
        if len(text.split()) < MIN_WORDS:
            print(f"[loader] skipped (too little text, maybe a scanned PDF): {path.name}")
            continue
        meta = _parse_header(text)
        doc_type = meta.get("type", default_type)
        if folder == "authoritative_cvs":
            doc_type = "cv"
        elif doc_type == "cv":
            print(f"[loader] skipped CV outside authoritative folder: {path}")
            continue
        if doc_type not in DOC_TYPES:
            print(f"[loader] unknown Type '{doc_type}' in {path.name}; using '{default_type}'")
            doc_type = default_type
        clean = " ".join(text.split())
        title, summary = meta.get("title"), meta.get("summary")
        auto = clean
        if doc_type == "cv" and cvparse.has_sections(text):
            # A real CV: "Name - Job title", and its Profile paragraph as the summary.
            title = title or cvparse.cv_title(text)
            auto = cvparse.profile_paragraph(text) or clean
        seen.add(doc_id)
        docs.append(Document(
            doc_id=doc_id,
            doc_type=doc_type,
            title=title or path.stem.replace("_", " "),
            summary=summary or (auto[:220] + ("..." if len(auto) > 220 else "")),
            text=text,
            path=path.resolve().as_posix() if folder == "authoritative_cvs" else path.relative_to(data_dir).as_posix(),
        ))
    return docs


def chunk_words(text: str, size: int = config.CHUNK_SIZE_WORDS,
                overlap: int = config.CHUNK_OVERLAP_WORDS) -> list[str]:
    words = text.split()
    if len(words) <= size:
        return [" ".join(words)]
    step = max(1, size - overlap)
    chunks = []
    for start in range(0, len(words), step):
        chunks.append(" ".join(words[start:start + size]))
        if start + size >= len(words):
            break
    return chunks
