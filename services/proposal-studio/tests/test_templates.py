"""tests/test_templates.py — Design-only template import, library, endpoints and template layout.

Synthetic PPTX files carry unique marker strings in text boxes, a grouped shape and a table cell;
after import none of them may remain anywhere. When the real OliveSoft templates are present in
templates/inbox/, every string of 4+ characters of each file is checked the same way.
"""
import io
import json
import os
import shutil
import sys
import tempfile
import threading
import time
import unittest
import urllib.error
import urllib.request
from pathlib import Path

STUDIO_DIR = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(STUDIO_DIR))
sys.path.insert(0, str(Path(__file__).resolve().parent))

TEST_TEMP_DIR = tempfile.mkdtemp(prefix="rfp_templates_test_")
os.environ["SAVES_DIR"] = str(Path(TEST_TEMP_DIR) / "saves")
os.environ["STUDIO_API_KEY"] = "secret-test-key-12345"
TEMPLATES_TEST_DIR = str(Path(TEST_TEMP_DIR) / "templates")

import server
import _runtime
import layout_engine
import template_importer as T

_runtime.isolate(server)

MARKERS = ["SECRET_TITLE_7F3A", "SECRET_BODY_91C2", "SECRET_FOOTER_55D0", "SECRET_GROUP_A1B2", "SECRET_CELL_C3D4"]


def make_pptx(n_slides=3) -> bytes:
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    from PIL import Image
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    logo = io.BytesIO()
    Image.new("RGB", (120, 120), (30, 120, 200)).save(logo, "PNG")
    for i in range(n_slides):
        s = prs.slides.add_slide(prs.slide_layouts[6])
        s.background.fill.solid()
        s.background.fill.fore_color.rgb = RGBColor(0x0B, 0x1E, 0x3F)
        bar = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, Inches(0.5))   # header band
        bar.fill.solid(); bar.fill.fore_color.rgb = RGBColor(0x22, 0x73, 0xF4)
        logo.seek(0)
        s.shapes.add_picture(logo, Inches(12.2), Inches(6.6), Inches(0.8), Inches(0.8))       # corner logo
        t = s.shapes.add_textbox(Inches(0.6), Inches(0.8), Inches(9), Inches(1))
        t.text_frame.text = f"{MARKERS[0]} slide {i}"
        t.text_frame.paragraphs[0].runs[0].font.size = Pt(40)
        b = s.shapes.add_textbox(Inches(0.6), Inches(2.2), Inches(11), Inches(3.5))
        b.text_frame.text = MARKERS[1]
        b.text_frame.paragraphs[0].runs[0].font.size = Pt(20)
        f = s.shapes.add_textbox(Inches(0.6), Inches(7.0), Inches(4), Inches(0.4))
        f.text_frame.text = MARKERS[2]
        g = s.shapes.add_group_shape()
        gs = g.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(10), Inches(1), Inches(2), Inches(1))
        gs.text_frame.text = MARKERS[3]
        tbl = s.shapes.add_table(2, 2, Inches(1), Inches(5.6), Inches(4), Inches(0.8)).table
        tbl.cell(1, 1).text = MARKERS[4]
    out = io.BytesIO()
    prs.save(out)
    return out.getvalue()


class TestImporter(unittest.TestCase):
    def setUp(self):
        self.dir = Path(tempfile.mkdtemp(prefix="rfp_tpl_lib_", dir=TEST_TEMP_DIR))
        self.lib = T.TemplateLibrary(self.dir)

    def tearDown(self):
        shutil.rmtree(self.dir, ignore_errors=True)

    def _all_text(self, tid):
        folder = self.dir / "library" / tid
        parts = [(folder / "template.json").read_text(encoding="utf-8")]
        parts += [p.read_text(encoding="utf-8", errors="ignore") for p in (folder / "assets").glob("*.svg")]
        return "\n".join(parts)

    def test_design_only_no_marker_anywhere(self):
        t = self.lib.import_bytes(make_pptx(), "Brand Test.pptx")
        blob = self._all_text(t["id"])
        for m in MARKERS:
            self.assertNotIn(m, blob)
        for role, els in t["designs"].items():
            self.assertTrue(els, role)
            for el in els:
                self.assertNotIn(el["type"], ("text", "bullets", "table", "card", "video"))
                self.assertTrue(el["locked"])
                self.assertLess(el["zIndex"], 100)

    def test_zones_and_theme_contrast(self):
        t = self.lib.import_bytes(make_pptx(), "brand.pptx")
        z = t["zones"]["content"]
        self.assertIsNotNone(z["title"])
        self.assertIsNotNone(z["body"])
        for box in (z["title"], z["body"]):
            self.assertTrue(0 <= box["x"] and box["x"] + box["w"] <= 1920 and 0 <= box["y"] and box["y"] + box["h"] <= 1080)
        self.assertGreaterEqual(T.contrast(t["theme"]["text"], t["theme"]["bg"]), 4.5)

    def test_ids_duplicates_and_one_or_two_slides(self):
        self.assertEqual(T.make_template_id("OliveSoft Commercial.pptx"), "olivesoft-commercial")
        a = self.lib.import_bytes(make_pptx(1), "Deck.pptx")
        b = self.lib.import_bytes(make_pptx(2), "Deck.pptx")
        self.assertEqual((a["id"], b["id"]), ("deck", "deck-2"))
        self.assertEqual(a["slideRoles"], {"title": 0, "content": 0, "thankyou": 0})
        self.assertEqual(b["slideRoles"], {"title": 0, "content": 1, "thankyou": 1})

    def test_rejections(self):
        with self.assertRaises(T.TemplateImportError):
            self.lib.import_bytes(b"not a zip", "bad.pptx")
        with self.assertRaises(T.TemplateImportError):
            self.lib.import_bytes(b"x", "notes.txt")
        with self.assertRaises(T.TemplateImportError):
            T.import_template(b"0" * (T.MAX_BYTES + 1), "big.pptx", lambda b, e: "")
        self.assertEqual(self.lib.ids(), [])        # nothing half-imported

    def test_rescan_default_delete_to_trash(self):
        (self.lib.inbox / "One.pptx").write_bytes(make_pptx())
        (self.lib.inbox / "broken.pptx").write_bytes(b"garbage")
        res = self.lib.rescan()
        self.assertEqual(res["imported"], ["one"])
        self.assertIn("broken.pptx", res["errors"])
        self.assertEqual(self.lib.get_default(), "one")
        self.assertEqual(self.lib.rescan()["imported"], [])     # already imported (same sha256)
        self.assertTrue(self.lib.delete("one"))
        self.assertIsNone(self.lib.get_default())
        self.assertEqual(len(list(self.lib.trash.iterdir())), 1)
        self.assertTrue((self.lib.inbox / "One.pptx").exists())  # the inbox file is never touched

    def test_asset_path_traversal(self):
        t = self.lib.import_bytes(make_pptx(), "brand.pptx")
        self.assertIsNone(self.lib.asset_path(t["id"], "../template.json"))
        self.assertIsNone(self.lib.asset_path("../../etc", "passwd"))

    def test_pdf_template_is_approximate_and_text_free(self):
        import pymupdf
        doc = pymupdf.open()
        for i in range(3):
            pg = doc.new_page(width=960, height=540)
            pg.draw_rect(pg.rect, fill=(0.05, 0.1, 0.3), color=None)
            pg.insert_text((60, 80), f"{MARKERS[0]} {i}", fontsize=28, color=(1, 1, 1))
        t = self.lib.import_bytes(doc.tobytes(), "brand.pdf")
        self.assertTrue(t["approximate"])
        png = next((self.dir / "library" / t["id"] / "assets").glob("*.png"))
        rendered = pymupdf.open(stream=png.read_bytes(), filetype="png")
        self.assertNotIn(MARKERS[0], rendered[0].get_text())

    @unittest.skipUnless(list((STUDIO_DIR / "templates" / "inbox").glob("*.pptx")), "real OliveSoft templates not present")
    def test_real_olivesoft_templates_contain_no_template_text(self):
        for f in sorted((STUDIO_DIR / "templates" / "inbox").glob("*.pptx")):
            data = f.read_bytes()
            t = self.lib.import_bytes(data, f.name)
            # The design parts carry no template text. Structured templates keep their fixed pages
            # ("À propos", contact) and placeholders under "structured" on purpose.
            designs = json.dumps({k: v for k, v in t.items() if k != "structured"}, ensure_ascii=False)
            blob = "\n".join([designs] + [p.read_text(encoding="utf-8", errors="ignore")
                                          for p in (self.dir / "library" / t["id"] / "assets").glob("*.svg")])
            leaked = [s for s in T._strings_worth_checking(T.original_strings_pptx(data))
                      if f'"{s}"' in blob or f">{s}<" in blob]
            self.assertEqual(leaked, [], f.name)


class TestTemplateLayout(unittest.TestCase):
    def template(self, body=None):
        return {"id": "t1", "theme": {"bg": "#0b1e3f", "primary": "#2273f4", "secondary": "#00dddc",
                                      "text": "#ffffff", "background": {"type": "solid", "color": "#0b1e3f"}},
                "designs": {"title": [{"type": "shape", "x": 0, "y": 0, "width": 1920, "height": 60, "zIndex": 1}],
                            "content": [{"type": "image", "src": "/x.png", "x": 1800, "y": 980, "width": 80, "height": 80, "zIndex": 1}],
                            "thankyou": []},
                "zones": {"title": {"title": {"x": 80, "y": 300, "w": 1200, "h": 120}, "subtitle": None},
                          "content": {"title": {"x": 80, "y": 70, "w": 600, "h": 60}, "body": body},
                          "thankyou": {"title": None, "qr": None}}}

    def spec(self, slides):
        return {"name": "Deck", "slides": slides}

    def test_without_template_unchanged(self):
        spec = json.loads((STUDIO_DIR / "samples" / "software.json").read_text(encoding="utf-8"))
        a, _ = layout_engine.layout_deck(spec)
        self.assertNotIn("template", a)
        self.assertFalse(any(e.get("locked") for s in a["slides"] for e in s["elements"]))

    def test_decor_below_content_and_text_in_zones(self):
        body = {"x": 80, "y": 180, "w": 1600, "h": 760}
        p, _ = layout_engine.layout_deck(self.spec([{"type": "content", "title": "T", "blocks": [{"kind": "text", "text": "x"}]}]),
                                         template=self.template(body))
        content = p["slides"][1]["elements"]
        decor = [e for e in content if e.get("locked")]
        others = [e for e in content if not e.get("locked")]
        self.assertTrue(decor and others)
        self.assertLess(max(e["zIndex"] for e in decor), min(e["zIndex"] for e in others))
        title = next(e for e in others if e["id"].startswith("el_title"))
        self.assertEqual((title["x"], title["y"]), (80, 70))
        self.assertEqual(p["slides"][0]["elements"][0]["type"], "shape")   # title design first
        self.assertEqual(p["template"], "t1")

    def test_narrow_body_zone_splits_by_its_height(self):
        # 6 table rows + 3 bullets need about 530 px: one slide in the default 760 px area,
        # two slides in a 520 px template body zone
        rows = [[f"r{i}", "x", "y"] for i in range(6)]
        slide = {"type": "content", "title": "Long", "blocks": [{"kind": "table", "columns": ["a", "b", "c"], "rows": rows},
                                                                 {"kind": "bullets", "items": [f"point {i}" for i in range(3)]}]}
        default, _ = layout_engine.layout_deck(self.spec([slide]))
        small, _ = layout_engine.layout_deck(self.spec([slide]), template=self.template({"x": 80, "y": 180, "w": 1300, "h": 520}))
        self.assertEqual(len(default["slides"]), 3)
        self.assertEqual(len(small["slides"]), 4)
        for s in small["slides"]:
            for e in s["elements"]:
                self.assertLessEqual(e["x"] + e.get("width", 0), 1920)

    def test_contrast_guard(self):
        t = self.template()
        t["theme"]["text"] = "#0c1f40"   # nearly the background color
        warnings = []
        theme = layout_engine.merge_template_theme(t, {}, warnings)
        self.assertEqual(theme["text"], "#FFFFFF")
        self.assertTrue(warnings)


class TestTemplateEndpoints(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        # Own template folder for this class (other test modules share the process environment)
        cls._prev_templates = os.environ.get("TEMPLATES_DIR")
        os.environ["TEMPLATES_DIR"] = TEMPLATES_TEST_DIR
        lib = server.get_template_library()
        (lib.inbox / "Brand.pptx").write_bytes(make_pptx())
        lib.rescan()   # auto-import is disabled in tests (TEMPLATES_AUTOIMPORT=0)
        cls.server = server.ThreadingHTTPServer(("127.0.0.1", 0), server.StudioHandler)
        cls.server.saves_dir = Path(os.environ["SAVES_DIR"])
        cls.base = f"http://127.0.0.1:{cls.server.server_address[1]}"
        threading.Thread(target=cls.server.serve_forever, daemon=True).start()
        time.sleep(0.1)

    @classmethod
    def tearDownClass(cls):
        cls.server.shutdown()
        cls.server.server_close()
        if cls._prev_templates is None:
            os.environ.pop("TEMPLATES_DIR", None)
        else:
            os.environ["TEMPLATES_DIR"] = cls._prev_templates
        shutil.rmtree(TEST_TEMP_DIR, ignore_errors=True)

    def setUp(self):
        os.environ["TEMPLATES_DIR"] = TEMPLATES_TEST_DIR

    def call(self, method, path, body=None, headers=None):
        req = urllib.request.Request(self.base + path, data=body, method=method, headers=headers or {})
        try:
            with urllib.request.urlopen(req) as r:
                return r.status, json.loads(r.read() or b"{}")
        except urllib.error.HTTPError as e:
            raw = e.read()
            try:
                return e.code, json.loads(raw)
            except Exception:
                return e.code, {}

    def test_list_get_default(self):
        code, res = self.call("GET", "/api/templates")
        self.assertEqual(code, 200)
        ids = [t["id"] for t in res["templates"]]
        self.assertIn("brand", ids)
        self.assertEqual(res["default"], "brand")
        self.assertEqual(self.call("GET", "/api/templates/brand")[1]["id"], "brand")
        self.assertEqual(self.call("GET", "/api/templates/nope")[0], 404)
        code, _ = self.call("PUT", "/api/templates/default", json.dumps({"id": "nope"}).encode())
        self.assertEqual(code, 400)

    def test_deck_with_template_and_unknown_template(self):
        spec = {"name": "D", "template": "brand", "slides": [{"type": "content", "title": "T", "blocks": [{"kind": "text", "text": "x"}]}]}
        h = {"Content-Type": "application/json", "X-API-Key": "secret-test-key-12345"}
        code, res = self.call("POST", "/api/decks", json.dumps(spec).encode(), h)
        self.assertEqual(code, 200, res)
        spec["template"] = "unknown"
        code, res = self.call("POST", "/api/decks", json.dumps(spec).encode(), h)
        self.assertEqual(code, 400)
        self.assertIn("unknown template 'unknown'", res["errors"][0])

    def test_upload_and_traversal(self):
        boundary = "XyZ"
        body = (f"--{boundary}\r\nContent-Disposition: form-data; name=\"file\"; filename=\"Second.pptx\"\r\n"
                "Content-Type: application/octet-stream\r\n\r\n").encode() + make_pptx() + f"\r\n--{boundary}--\r\n".encode()
        code, res = self.call("POST", "/api/templates", body, {"Content-Type": f"multipart/form-data; boundary={boundary}"})
        self.assertEqual(code, 200, res)
        bad = (f"--{boundary}\r\nContent-Disposition: form-data; name=\"file\"; filename=\"x.pptx\"\r\n\r\n").encode() + b"garbage" + f"\r\n--{boundary}--\r\n".encode()
        code, res = self.call("POST", "/api/templates", bad, {"Content-Type": f"multipart/form-data; boundary={boundary}"})
        self.assertEqual(code, 400)
        self.assertEqual(self.call("GET", "/templates/assets/brand/..%2Ftemplate.json")[0], 404)


if __name__ == "__main__":
    unittest.main()


INBOX = STUDIO_DIR / "templates" / "inbox"
REAL = sorted(INBOX.glob("template *.pptx"))


@unittest.skipUnless(REAL, "structured OliveSoft templates not present")
class TestStructuredTemplates(unittest.TestCase):
    """Cover / Sommaire / fixed pages / content / closing templates (placeholder words)."""

    def setUp(self):
        self.dir = Path(tempfile.mkdtemp(prefix="rfp_tpl_struct_", dir=TEST_TEMP_DIR))
        self.lib = T.TemplateLibrary(self.dir)

    def deck(self, t, slides=None, agenda=None):
        spec = {"name": "Proposition Portail", "client": "Ministère Démo", "date": "20 octobre 2026",
                "slides": slides or [{"type": "content", "title": f"Partie {i}", "blocks": [{"kind": "text", "text": "x"}]} for i in range(3)]}
        if agenda:
            spec["agenda"] = agenda
        return layout_engine.layout_deck(spec, template=t)[0]

    def test_roles_order_and_fixed_pages(self):
        for f in REAL:
            t = self.lib.import_bytes(f.read_bytes(), f.name)
            self.assertIn("structured", t, f.name)
            p = self.deck(t, agenda=["Synthèse", "Solution", "Planning"])
            labels = [s["label"] for s in p["slides"]]
            self.assertEqual(labels[:3], ["Titre & Couverture", "Sommaire", "À propos"], f.name)
            self.assertEqual(labels[-1], "Conclusion & Contact")
            self.assertEqual(len(p["slides"]), 3 + 3 + 1)
            texts = lambda s: [e.get("content") for e in s["elements"] if e.get("type") == "text"]
            # cover placeholders filled, no placeholder word left anywhere
            self.assertIn("Proposition Portail", texts(p["slides"][0]))
            all_text = " ".join(str(c) for s in p["slides"] for c in texts(s))
            for word in ("place1", "Titre de la diapositive", "Contenu"):
                self.assertNotIn(word, all_text, f.name)
            # agenda: 3 entries, unused slots removed
            agenda = p["slides"][1]["elements"]
            self.assertEqual(sorted(e["content"] for e in agenda if e.get("role") == "slot-label"),
                             ["Planning", "Solution", "Synthèse"])
            self.assertFalse(any(e.get("slot", 0) >= 3 for e in agenda))
            # À propos and closing copied from the template, with their text
            src = T.original_strings_pptx(f.read_bytes())
            self.assertTrue(any("OliveSoft" in str(c) for c in texts(p["slides"][2])))
            self.assertTrue(any("olivesoft.fr" in str(c) for c in texts(p["slides"][-1])))
            # real page numbers
            pages = [e["content"] for s in p["slides"][3:6] for e in s["elements"] if e.get("role") == "page-number"]
            self.assertEqual(pages, ["4", "5", "6"])


class TestNoClassDiagram(unittest.TestCase):
    def test_agents_never_produce_a_class_slide(self):
        import subprocess
        out = subprocess.run(["node", "-e", "const m=require('./agents/lib/merge.js');console.log(JSON.stringify(m.SLIDE_ORDER))"],
                             cwd=STUDIO_DIR, capture_output=True, text=True)
        self.assertNotIn("s3c", out.stdout)
        self.assertIn("s3b", out.stdout)
