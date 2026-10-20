"""tests/test_pptx.py — End-to-end tests for Phase 4 Native PowerPoint (.pptx) object export.

Asserts:
- Slide count matches expected slides
- Real native table objects for tables (rows and columns match)
- Text readable back from native text boxes and bullets
- Speaker notes present on slides with notes
- No slide consists of a single full-size picture (not screenshot-based)
- No shape lies outside the slide dimensions
- 40-row split table decks export correctly
- Deck with SVG images exports correctly
- Checks whether LibreOffice is installed (and reports status)
"""

import os
import sys
import shutil
import tempfile
import unittest
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches

STUDIO_DIR = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(STUDIO_DIR))

import layout_engine
import pptx_exporter


class TestNativePPTXExport(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.temp_dir = Path(tempfile.mkdtemp(prefix="rfp_pptx_test_"))
        cls.samples_dir = STUDIO_DIR / "samples"
        cls.libreoffice_path = shutil.which("libreoffice") or shutil.which("soffice")

    @classmethod
    def tearDownClass(cls):
        shutil.rmtree(cls.temp_dir, ignore_errors=True)

    def test_libreoffice_status(self):
        """Check if LibreOffice is installed in the test environment."""
        if self.libreoffice_path:
            print(f"[LibreOffice]: Found at {self.libreoffice_path}")
        else:
            print("[LibreOffice]: LibreOffice is not installed in this environment; PDF conversion comparison skipped.")

    def test_custom_shape_corner_radius_is_preserved(self):
        """A canvas corner radius maps to the matching native rounded-rect adjustment."""
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        # 48px is one sixth of the shortest canvas side (288px).
        element = {
            "type": "shape", "shapeType": "rectangle",
            "width": 576, "height": 288, "borderRadius": 48,
            "fill": "#6366f1", "stroke": "transparent", "strokeWidth": 0,
        }
        pptx_exporter.render_native_shape(
            slide, element, {}, Inches(1), Inches(1), Inches(4), Inches(2)
        )

        shape = slide.shapes[-1]
        self.assertAlmostEqual(shape.adjustments[0], 1 / 6, places=4)

    def test_imported_editable_path_honors_outline_and_corner_edits(self):
        """Imported path styling must override the source's no-outline metadata."""
        svg = pptx_exporter.editable_path_svg({
            "shapeType": "path", "pathWidth": 300, "pathHeight": 120,
            "borderRadius": 18, "fill": "#22c55e", "stroke": "#ffffff",
            "strokeWidth": 5, "strokeStyle": "dashed", "pathStyleOverride": True,
            "pathData": [{"d": "M0 0 H300 V120 H0 Z", "noFill": False, "noStroke": True}],
        })
        self.assertIn('stroke="#ffffff"', svg)
        self.assertIn('stroke-width="5"', svg)
        self.assertIn('stroke-dasharray="20 15"', svg)
        self.assertIn('rx="18"', svg)

    def test_export_sample_decks_to_native_pptx(self):
        """Export software.json, non_software.json, and thin.json to PPTX and verify native structure."""
        samples = ["software.json", "non_software.json", "thin.json"]

        for s_name in samples:
            sample_file = self.samples_dir / s_name
            self.assertTrue(sample_file.exists(), f"Sample {s_name} missing")
            import json
            deck_spec = json.loads(sample_file.read_text(encoding="utf-8"))

            # Run layout engine to produce full deck project with elements
            laid_out_deck, _ = layout_engine.layout_deck(deck_spec)

            out_pptx = self.temp_dir / f"{s_name}.pptx"
            ok, warnings = pptx_exporter.export_deck_to_pptx(laid_out_deck, out_pptx)
            self.assertTrue(ok)
            self.assertTrue(out_pptx.exists(), f"Failed to generate {out_pptx}")

            # Reopen generated presentation with python-pptx
            prs = Presentation(str(out_pptx))

            # 1. Slide count must match layout engine slides
            self.assertEqual(len(prs.slides), len(laid_out_deck["slides"]))

            # Dimensions must be 16:9 widescreen (13.333 x 7.5 in)
            self.assertAlmostEqual(prs.slide_width.inches, 13.333333, places=2)
            self.assertAlmostEqual(prs.slide_height.inches, 7.5, places=2)

            # Inspect each slide
            for idx, slide in enumerate(prs.slides):
                source_slide = laid_out_deck["slides"][idx]

                # 2. Check no slide consists of ONLY one full-size picture
                pictures = [s for s in slide.shapes if s.has_text_frame is False and s.has_table is False and getattr(s, "image", None) is not None]
                if len(slide.shapes) == 1 and pictures:
                    pic = pictures[0]
                    is_full_slide = (pic.width.inches >= 13.0 and pic.height.inches >= 7.0)
                    self.assertFalse(is_full_slide, f"Slide #{idx+1} in {s_name} is a full-slide screenshot fallback!")

                # 3. Check shape bounding box: no shape should lie completely outside the slide
                for shape in slide.shapes:
                    self.assertGreaterEqual(shape.left.inches + shape.width.inches, 0.0)
                    self.assertGreaterEqual(shape.top.inches + shape.height.inches, 0.0)
                    self.assertLessEqual(shape.left.inches, 13.5)
                    self.assertLessEqual(shape.top.inches, 8.0)

                # 4. If source slide had a table block, slide must contain a real native table shape
                has_table_block = any(el.get("type") == "table" for el in source_slide.get("elements", []))
                if has_table_block:
                    table_shapes = [s for s in slide.shapes if s.has_table]
                    self.assertGreater(len(table_shapes), 0, f"Slide #{idx+1} in {s_name} missing native table")
                    tbl = table_shapes[0].table
                    # Verify table has rows and columns
                    self.assertGreaterEqual(len(tbl.rows), 2)
                    self.assertGreaterEqual(len(tbl.columns), 2)
                    # Verify text is readable from cells
                    cell_text = tbl.cell(0, 0).text_frame.text
                    self.assertTrue(len(cell_text) > 0, "Table header cell must contain text")

                # 5. Verify text is readable from textboxes / bullets
                text_shapes = [s for s in slide.shapes if s.has_text_frame]
                self.assertGreater(len(text_shapes), 0, f"Slide #{idx+1} must contain readable native text")
                all_text = " ".join(s.text_frame.text for s in text_shapes)
                self.assertTrue(len(all_text.strip()) > 0)

                # 6. Verify notes if present on source slide
                if source_slide.get("notes"):
                    notes_slide = slide.notes_slide
                    self.assertIn(source_slide["notes"], notes_slide.notes_text_frame.text)

    def test_40_row_table_pptx_export(self):
        """Test exporting a 40-row split table deck to native PPTX."""
        rows = [[f"Critère {i}", f"Exigence {i}", f"Réponse {i}"] for i in range(1, 41)]
        spec = {
            "name": "Matrice Longue 40 Lignes",
            "client": "Client Grand Format",
            "slides": [
                {
                    "type": "content",
                    "title": "Matrice Complète",
                    "blocks": [
                        {
                            "kind": "table",
                            "columns": ["Réf", "Exigence", "Engagement"],
                            "rows": rows
                        }
                    ]
                }
            ]
        }
        laid_out, _ = layout_engine.layout_deck(spec)
        out_pptx = self.temp_dir / "large_table.pptx"
        ok, warnings = pptx_exporter.export_deck_to_pptx(laid_out, out_pptx)
        self.assertTrue(ok)
        self.assertTrue(out_pptx.exists())

        prs = Presentation(str(out_pptx))
        # Title + 4 table continuation slides + Thank you = 6 slides
        self.assertEqual(len(prs.slides), 6)

        # Inspect table slides
        table_slide_count = 0
        for s in prs.slides:
            tables = [shape for shape in s.shapes if shape.has_table]
            if tables:
                table_slide_count += 1
                tbl = tables[0].table
                self.assertEqual(len(tbl.columns), 3)
                self.assertEqual(tbl.cell(0, 0).text_frame.text, "Réf")
                self.assertEqual(tbl.cell(0, 1).text_frame.text, "Exigence")
                self.assertEqual(tbl.cell(0, 2).text_frame.text, "Engagement")
        self.assertEqual(table_slide_count, 4)

    def test_svg_image_pptx_export(self):
        """Test exporting a deck with an SVG image to native PPTX."""
        svg_content = """<svg xmlns="http://www.w3.org/2000/svg" width="300" height="200" viewBox="0 0 300 200">
  <rect width="300" height="200" fill="#6366f1" rx="20"/>
  <circle cx="150" cy="100" r="50" fill="#ffffff"/>
  <text x="150" y="110" font-size="20" text-anchor="middle" fill="#6366f1">TEST SVG</text>
</svg>"""
        import base64
        svg_dataurl = "data:image/svg+xml;base64," + base64.b64encode(svg_content.encode("utf-8")).decode("utf-8")

        spec = {
            "name": "Deck Test SVG",
            "client": "Client Vector",
            "slides": [
                {
                    "type": "content",
                    "title": "Illustration Vectorielle",
                    "blocks": [
                        {
                            "kind": "image",
                            "src": svg_dataurl,
                            "caption": "Schéma technique en SVG"
                        }
                    ]
                }
            ]
        }
        laid_out, _ = layout_engine.layout_deck(spec)
        out_pptx = self.temp_dir / "svg_deck.pptx"
        ok, warnings = pptx_exporter.export_deck_to_pptx(laid_out, out_pptx)
        self.assertTrue(ok)
        self.assertTrue(out_pptx.exists())

        prs = Presentation(str(out_pptx))
        # Slide 1 is content slide with the image
        content_slide = prs.slides[1]
        pictures = [shape for shape in content_slide.shapes if getattr(shape, "image", None) is not None]
        self.assertEqual(len(pictures), 1, "Rasterized SVG must be embedded as a picture in PPTX")


if __name__ == "__main__":
    unittest.main()
