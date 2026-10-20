"""tests/test_export_media.py — Uploaded images and SVG diagrams must appear in PDF and PPTX exports.

Regression for two bugs: the PDF was rendered from a file:// page, so /uploads/... images resolved
to file:///uploads/... and vanished; and SVG rasterization was disabled in the PPTX exporter, so
diagrams and SVG template decor vanished from PPTX files.
"""
import base64
import io
import json
import os
import subprocess
import sys
import tempfile
import threading
import unittest
import urllib.request
from pathlib import Path

STUDIO_DIR = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(STUDIO_DIR))
sys.path.insert(0, str(Path(__file__).resolve().parent))
os.environ["STUDIO_API_KEY"] = "secret-test-key-12345"

import server
import _runtime

_runtime.isolate(server)
KEY = "secret-test-key-12345"
SVG = ('<svg xmlns="http://www.w3.org/2000/svg" width="800" height="400" viewBox="0 0 800 400">'
       '<rect width="800" height="400" fill="#0000ff"/></svg>')


class TestExportMedia(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.tmp = Path(tempfile.mkdtemp(prefix="rfp_export_media_"))
        cls.srv = server.ThreadingHTTPServer(("127.0.0.1", 0), server.StudioHandler)
        cls.srv.saves_dir = cls.tmp / "saves"
        cls.srv.saves_dir.mkdir()
        cls.base = f"http://127.0.0.1:{cls.srv.server_address[1]}"
        threading.Thread(target=cls.srv.serve_forever, daemon=True).start()

    @classmethod
    def tearDownClass(cls):
        cls.srv.shutdown()
        cls.srv.server_close()

    def upload(self, name, data):
        b = "B0"
        body = (f"--{b}\r\nContent-Disposition: form-data; name=\"file\"; filename=\"{name}\"\r\n\r\n").encode() + data + f"\r\n--{b}--\r\n".encode()
        req = urllib.request.Request(self.base + "/api/upload", data=body, headers={"Content-Type": f"multipart/form-data; boundary={b}"})
        return json.load(urllib.request.urlopen(req))["url"]

    @_runtime.slow
    def test_uploaded_png_svg_upload_and_svg_data_url_in_pdf_and_pptx(self):
        os.environ["STUDIO_API_KEY"] = KEY
        from PIL import Image
        buf = io.BytesIO()
        Image.new("RGB", (400, 300), (255, 0, 0)).save(buf, "PNG")
        png_url = self.upload("photo.png", buf.getvalue())
        svg_url = self.upload("diagram.svg", SVG.encode())
        data_url = "data:image/svg+xml;base64," + base64.b64encode(SVG.replace("#0000ff", "#00ff00").encode()).decode()
        spec = {"name": "Media", "theme": {"background": "#ffffff"}, "slides": [
            {"type": "content", "title": "Photo", "blocks": [{"kind": "image", "src": png_url}]},
            {"type": "content", "title": "SVG upload", "blocks": [{"kind": "image", "src": svg_url}]},
            {"type": "content", "title": "SVG data", "blocks": [{"kind": "image", "src": data_url}]}]}
        req = urllib.request.Request(self.base + "/api/decks", data=json.dumps(spec).encode(),
                                     headers={"Content-Type": "application/json", "X-API-Key": KEY})
        deck_id = json.load(urllib.request.urlopen(req))["id"]

        files = {}
        for fmt in ("pdf", "pptx"):
            out = self.tmp / f"deck.{fmt}"
            r = subprocess.run([sys.executable, str(STUDIO_DIR / "tools" / "export_deck.py"), deck_id, fmt,
                                "--base-url", self.base, "--out", str(out), "--force"], capture_output=True, text=True, timeout=180)
            self.assertTrue(out.exists(), r.stdout + r.stderr)
            files[fmt] = out

        import pymupdf
        pdf = pymupdf.open(files["pdf"])

        def count(page, rgb):
            px = pdf[page].get_pixmap(dpi=30).samples
            return sum(1 for i in range(0, len(px), 3)
                       if all(abs(px[i + k] - rgb[k]) < 40 for k in range(3)))
        self.assertGreater(count(1, (255, 0, 0)), 500, "uploaded PNG missing from the PDF")
        self.assertGreater(count(2, (0, 0, 255)), 500, "uploaded SVG missing from the PDF")
        self.assertGreater(count(3, (0, 255, 0)), 500, "SVG data URL missing from the PDF")

        from pptx import Presentation
        prs = Presentation(files["pptx"])
        pictures = [sum(1 for s in sl.shapes if s.shape_type == 13) for sl in prs.slides]
        self.assertEqual(pictures[1:4], [1, 1, 1], f"pictures per slide: {pictures}")

    @_runtime.slow
    def test_text_hyperlink_in_pdf_and_pptx(self):
        os.environ["STUDIO_API_KEY"] = KEY
        spec = {"name": "Links", "slides": [{"type": "content", "title": "Lien", "blocks": [{"kind": "text", "text": "Voir le site"}]}]}
        req = urllib.request.Request(self.base + "/api/decks", data=json.dumps(spec).encode(),
                                     headers={"Content-Type": "application/json", "X-API-Key": KEY})
        deck_id = json.load(urllib.request.urlopen(req))["id"]
        save = self.srv.saves_dir / f"{deck_id}.json"
        deck = json.loads(save.read_text(encoding="utf-8"))
        el = next(e for e in deck["slides"][1]["elements"] if e.get("content") == "Voir le site")
        el["link"] = "https://olivesoft.fr/"
        save.write_text(json.dumps(deck), encoding="utf-8")
        files = {}
        for fmt in ("pdf", "pptx"):
            out = self.tmp / f"links.{fmt}"
            r = subprocess.run([sys.executable, str(STUDIO_DIR / "tools" / "export_deck.py"), deck_id, fmt,
                                "--base-url", self.base, "--out", str(out), "--force"], capture_output=True, text=True, timeout=180)
            self.assertTrue(out.exists(), r.stdout + r.stderr)
            files[fmt] = out
        import pymupdf
        uris = [l.get("uri") for page in pymupdf.open(files["pdf"]) for l in page.get_links()]
        self.assertIn("https://olivesoft.fr/", uris)
        from pptx import Presentation
        addrs = [r.hyperlink.address for sl in Presentation(files["pptx"]).slides for sh in sl.shapes if sh.has_text_frame
                 for p in sh.text_frame.paragraphs for r in p.runs]
        self.assertIn("https://olivesoft.fr/", addrs)


if __name__ == "__main__":
    unittest.main()
