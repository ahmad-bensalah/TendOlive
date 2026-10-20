"""pptx_extractor.py — Import a PPTX template: theme (colors, fonts, background) and slides.

extract_theme(pptx_bytes, store_media=None) returns:
    {
      "colors":     { "dark1": "#...", "accent1": "#...", ... }   theme color scheme
      "fonts":      { "heading": {name, stack, weight}, "body": {...} }
      "background": { "type": "solid"|"gradient"|"image", ... }   project background
      "palette":    { bg, surface, text, text_muted, primary, secondary, accent, accent_2, accent_3 }
      "slides":     [ { type, label, elements: [...] } ]          editor slides (1920x1080)
      "warnings":   [ "2 charts skipped (not supported)", ... ]
    }

The slide import reads the slide XML directly so that it can follow PowerPoint's real
geometry: nested group coordinate systems, rotation and flips, placeholder positions
inherited from the layout and master, text styles inherited from the master, theme colors
with their modifiers, image crops, and the items placed on the master and layout slides.

Shapes become editable editor elements where the editor has an equivalent (rectangles,
rounded rectangles, ellipses, preset paths and freeforms, text, bullets, tables, images,
videos). Gradients and vector features without an editable equivalent remain SVG images
so their visual fidelity is preserved.

store_media(blob, ext) -> url saves an image or video and returns its URL. The server
passes a function that writes into uploads/ (deduplicated); without it, media is
returned as data URLs.
"""
import base64
import colorsys
import io
import math
import re
import zipfile
from collections import Counter
from typing import Any, Callable, Dict, List, Optional, Tuple

from lxml import etree

# ── namespaces ─────────────────────────────────────────────────────────
_NSMAP = {
    "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
    "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
    "p": "http://schemas.openxmlformats.org/presentationml/2006/main",
    "rel": "http://schemas.openxmlformats.org/package/2006/relationships",
    "mc": "http://schemas.openxmlformats.org/markup-compatibility/2006",
    "asvg": "http://schemas.microsoft.com/office/drawing/2016/SVG/main",
    "p14": "http://schemas.microsoft.com/office/powerpoint/2010/main",
}

CANVAS_W, CANVAS_H = 1920, 1080
EMU_PER_PT = 12700


def _tag(ns, local):
    return f"{{{_NSMAP[ns]}}}{local}"


def _q(path: str) -> str:
    """'a:spPr/a:solidFill' -> Clark notation path for lxml find()."""
    return "/".join(_tag(*part.split(":")) if ":" in part else part for part in path.split("/"))


def _find(el, path):
    return el.find(_q(path)) if el is not None else None


def _findall(el, path):
    return el.findall(_q(path)) if el is not None else []


def _local(el) -> str:
    return etree.QName(el).localname


def _int(v, default=0) -> int:
    try:
        return int(v)
    except (TypeError, ValueError):
        return default


# ── font helpers ───────────────────────────────────────────────────────
_FONT_FALLBACKS = {
    "calibri":           ("'Carlito', 'Calibri', sans-serif", 400),
    "arial":             ("'Arial', 'Helvetica', sans-serif", 400),
    "arial black":       ("'Arial Black', 'Gadget', sans-serif", 900),
    "times new roman":   ("'Times New Roman', 'Times', serif", 400),
    "georgia":           ("'Georgia', serif", 400),
    "verdana":           ("'Verdana', 'Geneva', sans-serif", 400),
    "tahoma":            ("'Tahoma', 'Geneva', sans-serif", 400),
    "trebuchet ms":      ("'Trebuchet MS', 'Lucida Grande', sans-serif", 400),
    "century gothic":    ("'Century Gothic', 'CenturyGothic', sans-serif", 400),
    "segoe ui":          ("'Segoe UI', 'Roboto', sans-serif", 400),
    "tt norms pro":      ("'Montserrat', 'Plus Jakarta Sans', sans-serif", 700),
    "gotham":            ("'Gotham', 'Montserrat', sans-serif", 400),
    "proxima nova":      ("'Proxima Nova', 'Montserrat', sans-serif", 400),
    "arimo":             ("'Arimo', 'Arial', sans-serif", 400),
    "canva sans":        ("'Canva Sans', 'Inter', sans-serif", 400),
    "montserrat":        ("'Montserrat', sans-serif", 600),
    "poppins":           ("'Poppins', sans-serif", 400),
    "roboto":            ("'Roboto', sans-serif", 400),
    "inter":             ("'Inter', sans-serif", 400),
    "open sans":         ("'Open Sans', sans-serif", 400),
    "lato":              ("'Lato', sans-serif", 400),
    "raleway":           ("'Raleway', sans-serif", 400),
    "plus jakarta sans": ("'Plus Jakarta Sans', sans-serif", 600),
    "outfit":            ("'Outfit', sans-serif", 400),
    "public sans":       ("'Public Sans', sans-serif", 400),
}

# Style words that exporters (Canva, Google Slides) append to the family name.
_STYLE_WORDS = {
    "thin": 100, "hairline": 100, "extralight": 200, "ultralight": 200, "light": 300,
    "regular": 400, "book": 400, "normal": 400, "medium": 500, "semibold": 600, "demibold": 600,
    "bold": 700, "extrabold": 800, "ultrabold": 800, "black": 900, "heavy": 900,
}


def _clean_font_name(name: str) -> Tuple[str, Optional[int], bool]:
    """'Arimo Bold' -> ('Arimo', 700, False); 'Public Sans Bold Italics' -> ('Public Sans', 700, True);
    'Calibri (MS) Bold' -> ('Calibri', 700, False)."""
    if not name:
        return "", None, False
    words = re.sub(r"\(.*?\)", " ", name).replace("-", " ").split()
    weight, italic = None, False
    while len(words) > 1:
        w = words[-1].lower()
        if w in ("italic", "italics", "oblique"):
            italic = True
        elif w in _STYLE_WORDS:
            weight = _STYLE_WORDS[w] if weight is None else max(weight, _STYLE_WORDS[w])
        else:
            break
        words.pop()
    return " ".join(words), weight, italic


def _resolve_font(name):
    """Given a font name, return (css_stack, weight). Falls back to sans-serif."""
    if not name:
        return ("sans-serif", 400)
    key = name.strip().lower()
    if key in _FONT_FALLBACKS:
        return _FONT_FALLBACKS[key]
    for k, v in _FONT_FALLBACKS.items():
        if k in key or key in k:
            return v
    return (f"'{name}', sans-serif", 400)


# ── colors ─────────────────────────────────────────────────────────────
_SCHEME_ORDER = ["dk1", "lt1", "dk2", "lt2", "accent1", "accent2", "accent3", "accent4",
                 "accent5", "accent6", "hlink", "folHlink"]
_SCHEME_LABELS = {
    "dk1": "dark1", "lt1": "light1", "dk2": "dark2", "lt2": "light2",
    "accent1": "accent1", "accent2": "accent2", "accent3": "accent3", "accent4": "accent4",
    "accent5": "accent5", "accent6": "accent6", "hlink": "hyperlink", "folHlink": "followed_hyperlink",
}
_PRESET_COLORS = {"black": "000000", "white": "ffffff", "red": "ff0000", "green": "008000",
                  "blue": "0000ff", "yellow": "ffff00", "gray": "808080", "grey": "808080"}
# accent1 of the stock Office themes: a template still using one of these has no brand colors in its theme
_OFFICE_DEFAULT_ACCENT1 = {"#4f81bd", "#4472c4", "#5b9bd5", "#156082", "#0f6fc6"}


def _hex(rgb: Tuple[float, float, float]) -> str:
    return "#" + "".join(f"{max(0, min(255, round(c))):02x}" for c in rgb)


def _rgb(hex_color: str) -> Tuple[int, int, int]:
    h = hex_color.lstrip("#")
    return int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16)


def _luminance(hex_color: str) -> float:
    def ch(c):
        c = c / 255
        return c / 12.92 if c <= 0.03928 else ((c + 0.055) / 1.055) ** 2.4
    r, g, b = _rgb(hex_color)
    return 0.2126 * ch(r) + 0.7152 * ch(g) + 0.0722 * ch(b)


def _contrast(c1: str, c2: str) -> float:
    l1, l2 = sorted((_luminance(c1), _luminance(c2)), reverse=True)
    return (l1 + 0.05) / (l2 + 0.05)


def _is_neutral(hex_color: str) -> bool:
    r, g, b = _rgb(hex_color)
    _, l, s = colorsys.rgb_to_hls(r / 255, g / 255, b / 255)
    return s < 0.15 or l < 0.08 or l > 0.94


class ColorResolver:
    """Resolves DrawingML color nodes (srgbClr, schemeClr, sysClr, prstClr) with modifiers."""

    def __init__(self, scheme: Dict[str, str], clr_map: Optional[Dict[str, str]] = None):
        self.scheme = scheme  # dk1 -> "#rrggbb"
        self.clr_map = clr_map or {"bg1": "lt1", "tx1": "dk1", "bg2": "lt2", "tx2": "dk2"}

    def resolve(self, fill_parent, ph_color: Optional[Tuple[str, float]] = None) -> Optional[Tuple[str, float]]:
        """fill_parent contains one color child (e.g. <a:solidFill>). Returns (hex, alpha) or None."""
        if fill_parent is None:
            return None
        for node in fill_parent:
            name = _local(node)
            base = None
            if name == "srgbClr":
                base = "#" + node.get("val", "000000").lower()
            elif name == "sysClr":
                base = "#" + node.get("lastClr", "000000").lower()
            elif name == "prstClr":
                base = "#" + _PRESET_COLORS.get(node.get("val", ""), "000000")
            elif name == "scrgbClr":
                base = _hex(tuple(_int(node.get(k)) / 100000 * 255 for k in ("r", "g", "b")))
            elif name == "schemeClr":
                val = node.get("val", "")
                if val == "phClr":
                    if ph_color is None:
                        continue
                    base = ph_color[0]
                else:
                    key = self.clr_map.get(val, val)
                    base = self.scheme.get(key)
            if base is None:
                continue
            return self._apply_modifiers(base, node)
        return None

    @staticmethod
    def _apply_modifiers(base: str, node) -> Tuple[str, float]:
        r, g, b = _rgb(base)
        alpha = 1.0
        h, l, s = colorsys.rgb_to_hls(r / 255, g / 255, b / 255)
        hls_changed = False
        for m in node:
            mname = _local(m)
            v = _int(m.get("val"), 100000) / 100000
            if mname == "alpha":
                alpha = v
            elif mname == "lumMod":
                l *= v
                hls_changed = True
            elif mname == "lumOff":
                l += v
                hls_changed = True
            elif mname == "satMod":
                s *= v
                hls_changed = True
            elif mname in ("tint", "shade"):
                if hls_changed:
                    r, g, b = (c * 255 for c in colorsys.hls_to_rgb(h, max(0, min(1, l)), max(0, min(1, s))))
                    hls_changed = False
                if mname == "tint":
                    r, g, b = (c + (255 - c) * (1 - v) for c in (r, g, b))
                else:
                    r, g, b = (c * v for c in (r, g, b))
                h, l, s = colorsys.rgb_to_hls(r / 255, g / 255, b / 255)
        if hls_changed:
            r, g, b = (c * 255 for c in colorsys.hls_to_rgb(h, max(0, min(1, l)), max(0, min(1, s))))
        return _hex((r, g, b)), alpha


def _css_color(color: Tuple[str, float]) -> str:
    hex_color, alpha = color
    if alpha >= 0.999:
        return hex_color
    r, g, b = _rgb(hex_color)
    return f"rgba({r}, {g}, {b}, {round(alpha, 3)})"


# ── geometry ───────────────────────────────────────────────────────────
class Transform:
    """Maps a shape's EMU coordinates (possibly inside nested groups) to slide EMU."""

    def __init__(self, groups=None):
        self.groups = groups or []  # outermost first

    def child(self, grp_xfrm) -> "Transform":
        return Transform(self.groups + [grp_xfrm])

    def apply(self, x: float, y: float, w: float, h: float, rot: float, flip_h: bool, flip_v: bool):
        """Returns (center_x, center_y, width, height, rotation_deg, flip_h, flip_v) in slide EMU."""
        cx, cy = x + w / 2, y + h / 2
        for g in reversed(self.groups):
            sx = g["ext"][0] / g["chExt"][0] if g["chExt"][0] else 1.0
            sy = g["ext"][1] / g["chExt"][1] if g["chExt"][1] else 1.0
            cx = g["off"][0] + (cx - g["chOff"][0]) * sx
            cy = g["off"][1] + (cy - g["chOff"][1]) * sy
            w, h = w * sx, h * sy
            gcx, gcy = g["off"][0] + g["ext"][0] / 2, g["off"][1] + g["ext"][1] / 2
            if g["flipH"]:
                cx = 2 * gcx - cx
                rot = -rot
                flip_h = not flip_h
            if g["flipV"]:
                cy = 2 * gcy - cy
                rot = -rot
                flip_v = not flip_v
            if g["rot"]:
                t = math.radians(g["rot"])
                dx, dy = cx - gcx, cy - gcy
                cx = gcx + dx * math.cos(t) - dy * math.sin(t)
                cy = gcy + dx * math.sin(t) + dy * math.cos(t)
                rot += g["rot"]
        return cx, cy, w, h, rot % 360, flip_h, flip_v


def _read_xfrm(xfrm) -> Optional[Dict[str, Any]]:
    if xfrm is None:
        return None
    off, ext = _find(xfrm, "a:off"), _find(xfrm, "a:ext")
    if off is None or ext is None:
        return None
    d = {
        "off": (_int(off.get("x")), _int(off.get("y"))),
        "ext": (_int(ext.get("cx")), _int(ext.get("cy"))),
        "rot": _int(xfrm.get("rot")) / 60000,
        "flipH": xfrm.get("flipH") in ("1", "true"),
        "flipV": xfrm.get("flipV") in ("1", "true"),
    }
    cho, che = _find(xfrm, "a:chOff"), _find(xfrm, "a:chExt")
    d["chOff"] = (_int(cho.get("x")), _int(cho.get("y"))) if cho is not None else d["off"]
    d["chExt"] = (_int(che.get("cx")), _int(che.get("cy"))) if che is not None else d["ext"]
    return d


# ── importer ───────────────────────────────────────────────────────────
class _Importer:
    def __init__(self, prs, store_media: Callable[[bytes, str], str], scheme: Dict[str, str],
                 major_font: str, minor_font: str):
        self.prs = prs
        self._store_media = store_media
        self.media_blobs: Dict[str, Tuple[bytes, str]] = {}
        self._ids = 0
        self.major_font = major_font
        self.minor_font = minor_font
        slide_cx, slide_cy = float(prs.slide_width), float(prs.slide_height)
        # Uniform scale (keeps 4:3 decks undistorted), content centered on the 1920x1080 canvas
        self.scale = min(CANVAS_W / slide_cx, CANVAS_H / slide_cy)
        self.dx = (CANVAS_W - slide_cx * self.scale) / 2
        self.dy = (CANVAS_H - slide_cy * self.scale) / 2
        self.px_per_pt = self.scale * EMU_PER_PT
        self.slide_emu = (slide_cx, slide_cy)
        self.scheme = scheme
        self.warnings: Counter = Counter()
        self.media_cache: Dict[Tuple[str, str], Optional[str]] = {}
        self.text_colors: Counter = Counter()
        self.fill_colors: Counter = Counter()
        self.run_fonts: Counter = Counter()
        self.big_run_fonts: Counter = Counter()

    def eid(self, prefix: str) -> str:
        self._ids += 1
        return f"el_{prefix}_{self._ids}"

    # ── media ──
    def store(self, blob: bytes, ext: str) -> str:
        url = self._store_media(blob, ext)
        self.media_blobs[url] = (blob, ext)
        return url

    def as_data_url(self, url: str) -> str:
        if url.startswith("data:") or url not in self.media_blobs:
            return url
        return _data_url_store(*self.media_blobs[url])

    def media_url(self, part, r_id: str, crop=None, mask: Optional[str] = None) -> Optional[str]:
        try:
            target = part.related_part(r_id)
        except Exception:
            return None
        blob = target.blob
        ext = str(target.partname.ext).lower()
        if ext in ("emf", "wmf"):
            self.warnings["EMF/WMF pictures skipped (not supported by browsers)"] += 1
            return None
        raster = ext in ("png", "jpg", "jpeg", "gif", "webp")
        if ext in ("tif", "tiff", "bmp") or (raster and (crop or mask)):
            blob, ext = self._raster(blob, ext, crop, mask)
            if blob is None:
                return None
        key = (str(target.partname), repr((crop, mask)) if ext != "svg" else "")
        if key not in self.media_cache:
            try:
                self.media_cache[key] = self.store(blob, "jpg" if ext == "jpeg" else ext)
            except ValueError as e:
                self.warnings[f"media files skipped ({e})"] += 1
                self.media_cache[key] = None
        return self.media_cache[key]

    @staticmethod
    def _raster(blob: bytes, ext: str, crop, mask=None) -> Tuple[Optional[bytes], str]:
        """Convert, crop (PPTX srcRect fractions) and mask ('ellipse' or 'roundRect') a raster image."""
        try:
            from PIL import Image, ImageDraw
            img = Image.open(io.BytesIO(blob))
            if crop:
                w, h = img.size
                l, t, r, b = (max(0.0, c) for c in crop)
                box = (round(w * l), round(h * t), round(w * (1 - r)), round(h * (1 - b)))
                if box[2] > box[0] and box[3] > box[1]:
                    img = img.crop(box)
            if mask:
                img = img.convert("RGBA")
                alpha = Image.new("L", img.size, 0)
                draw = ImageDraw.Draw(alpha)
                if mask == "ellipse":
                    draw.ellipse((0, 0, img.width - 1, img.height - 1), fill=255)
                else:
                    draw.rounded_rectangle((0, 0, img.width - 1, img.height - 1),
                                           radius=round(min(img.size) * 0.16667), fill=255)
                img.putalpha(Image.composite(img.getchannel("A"), alpha, alpha))
                ext = "png"
            out = io.BytesIO()
            if ext in ("jpg", "jpeg") and img.mode in ("RGB", "L"):
                img.save(out, format="JPEG", quality=92)
                return out.getvalue(), "jpg"
            img.save(out, format="PNG")
            return out.getvalue(), "png"
        except Exception:
            return None, ext

    def store_svg(self, svg: str) -> str:
        key = ("svg-inline", svg)
        if key not in self.media_cache:
            self.media_cache[key] = self.store(svg.encode("utf-8"), "svg")
        return self.media_cache[key]

    # ── geometry helpers ──
    def box(self, xfrm_d, transform: Transform):
        x, y = xfrm_d["off"]
        w, h = xfrm_d["ext"]
        cx, cy, w, h, rot, fh, fv = transform.apply(x, y, w, h, xfrm_d["rot"], xfrm_d["flipH"], xfrm_d["flipV"])
        pw, ph = w * self.scale, h * self.scale
        return {
            "x": round(self.dx + cx * self.scale - pw / 2),
            "y": round(self.dy + cy * self.scale - ph / 2),
            "width": max(1, round(pw)),
            "height": max(1, round(ph)),
            "rotation": round(rot, 2) if abs(rot) > 0.01 and abs(rot - 360) > 0.01 else 0,
        }, fh, fv

    def is_full_bleed(self, b) -> bool:
        return b["x"] <= 25 and b["y"] <= 25 and b["width"] >= CANVAS_W - 90 and b["height"] >= CANVAS_H - 50 and not b["rotation"]

    # ── slide walking ──
    def import_slide(self, slide) -> List[Dict[str, Any]]:
        layout = slide.slide_layout
        master = layout.slide_master
        ctx_master = _Ctx(self, master.part, master._element, None, None, "master")
        ctx_layout = _Ctx(self, layout.part, layout._element, None, ctx_master, "layout")
        ctx_slide = _Ctx(self, slide.part, slide._element, ctx_layout, ctx_master, "slide")

        elements: List[Dict[str, Any]] = []
        show_master = slide._element.get("showMasterSp", "1") not in ("0", "false")
        show_layout_master = layout._element.get("showMasterSp", "1") not in ("0", "false")
        if show_master and show_layout_master:
            elements += ctx_master.walk_tree(skip_placeholders=True)
        if show_master:
            elements += ctx_layout.walk_tree(skip_placeholders=True)
        elements += ctx_slide.walk_tree(skip_placeholders=False)
        # Items parked entirely outside the slide are invisible in PowerPoint: drop them
        visible = [e for e in elements if e["x"] < CANVAS_W and e["y"] < CANVAS_H
                   and e["x"] + e["width"] > 0 and e["y"] + e["height"] > 0]
        if len(visible) < len(elements):
            self.warnings["items outside the slide skipped"] += len(elements) - len(visible)
        return visible


class _Ctx:
    """One XML part (slide, layout or master) being imported."""

    def __init__(self, imp: _Importer, part, root, layout_ctx, master_ctx, kind):
        self.imp = imp
        self.part = part
        self.root = root
        self.layout_ctx = layout_ctx
        self.master_ctx = master_ctx
        self.kind = kind
        self._ph_index = None

    # placeholders: positions and text styles are inherited from layout, then master
    def placeholders(self):
        if self._ph_index is None:
            idx, typ = {}, {}
            for sp in self.root.iter(_tag("p", "sp"), _tag("p", "pic")):
                ph = sp.find(".//" + _tag("p", "ph"))
                if ph is None:
                    continue
                if ph.get("idx") is not None:
                    idx.setdefault(ph.get("idx"), sp)
                typ.setdefault(ph.get("type", "body"), sp)
            self._ph_index = (idx, typ)
        return self._ph_index

    def inherited_ph(self, ph) -> List[Any]:
        """The matching placeholder shapes in layout and master (closest first)."""
        found = []
        p_idx, p_type = ph.get("idx"), ph.get("type", "body")
        for ctx in (self.layout_ctx, self.master_ctx):
            if ctx is None:
                continue
            idx, typ = ctx.placeholders()
            match = idx.get(p_idx) if p_idx is not None else None
            if match is None:
                match = typ.get(p_type) or (typ.get("title") if p_type == "ctrTitle" else None) \
                    or (typ.get("body") if p_type in ("subTitle", "obj") else None)
            if match is not None:
                found.append(match)
        return found

    def walk_tree(self, skip_placeholders: bool) -> List[Dict[str, Any]]:
        sp_tree = self.root.find(".//" + _tag("p", "spTree"))
        out: List[Dict[str, Any]] = []
        if sp_tree is not None:
            self.walk(sp_tree, Transform(), out, skip_placeholders)
        return out

    def walk(self, parent, transform: Transform, out, skip_placeholders):
        for child in parent:
            name = _local(child)
            try:
                if name == "AlternateContent":
                    fallback = _find(child, "mc:Fallback")
                    choice = _find(child, "mc:Choice")
                    self.walk(fallback if fallback is not None else choice, transform, out, skip_placeholders)
                    continue
                nv = child.find(".//" + _tag("p", "cNvPr"))
                if nv is not None and nv.get("hidden") in ("1", "true"):
                    continue
                if name == "grpSp":
                    g = _read_xfrm(_find(child, "p:grpSpPr/a:xfrm"))
                    self.walk(child, transform.child(g) if g else transform, out, skip_placeholders)
                elif name in ("sp", "cxnSp", "pic"):
                    ph = child.find(".//" + _tag("p", "ph"))
                    if ph is not None and skip_placeholders:
                        continue
                    if name == "pic":
                        self.picture(child, ph, transform, out)
                    else:
                        self.shape(child, ph, transform, out)
                elif name == "graphicFrame":
                    self.graphic_frame(child, transform, out)
            except Exception:
                self.imp.warnings["items skipped because they could not be read"] += 1

    def xfrm_for(self, el, ph) -> Optional[Dict[str, Any]]:
        own = _read_xfrm(_find(el, "p:spPr/a:xfrm") if _local(el) != "graphicFrame" else _find(el, "p:xfrm"))
        if own is not None or ph is None:
            return own
        for inh in self.inherited_ph(ph):
            d = _read_xfrm(_find(inh, "p:spPr/a:xfrm"))
            if d is not None:
                return d
        return None

    # ── fills and lines ──
    def fill_of(self, sp_pr, style=None):
        """Returns ('none'|'solid'|'grad'|'blip'|None, data)."""
        if sp_pr is not None:
            for node in sp_pr:
                n = _local(node)
                if n == "noFill":
                    return "none", None
                if n == "solidFill":
                    return "solid", self.colors().resolve(node)
                if n == "gradFill":
                    return "grad", node
                if n == "blipFill":
                    return "blip", node
                if n == "grpFill":
                    return None, None
        if style is not None:
            ref = _find(style, "a:fillRef")
            if ref is not None and _int(ref.get("idx")) > 0:
                return "solid", self.colors().resolve(ref)
        return None, None

    def line_of(self, sp_pr, style=None):
        """Returns (css color or None, width_px, dash)."""
        ln = _find(sp_pr, "a:ln")
        color = None
        if ln is not None:
            if _find(ln, "a:noFill") is not None:
                return None, 0, "solid"
            solid = _find(ln, "a:solidFill")
            if solid is not None:
                color = self.colors().resolve(solid)
        if color is None and style is not None:
            ref = _find(style, "a:lnRef")
            if ref is not None and _int(ref.get("idx")) > 0 and (ln is None or _find(ln, "a:noFill") is None):
                color = self.colors().resolve(ref)
        if color is None:
            return None, 0, "solid"
        w_emu = _int(ln.get("w"), 12700) if ln is not None else 12700
        dash = _find(ln, "a:prstDash") if ln is not None else None
        dash_val = dash.get("val", "solid") if dash is not None else "solid"
        style_name = "dotted" if "dot" in dash_val.lower() and "dash" not in dash_val.lower() else \
            ("dashed" if dash_val != "solid" else "solid")
        return _css_color(color), max(1, round(w_emu * self.imp.scale)), style_name

    def colors(self) -> ColorResolver:
        return self.imp.color_resolver

    # ── shapes ──
    def shape(self, sp, ph, transform: Transform, out):
        imp = self.imp
        x = self.xfrm_for(sp, ph)
        if x is None:
            return
        b, flip_h, flip_v = imp.box(x, transform)
        sp_pr = _find(sp, "p:spPr")
        style = _find(sp, "p:style")
        geom = _find(sp_pr, "a:prstGeom")
        prst = geom.get("prst") if geom is not None else None
        cust = _find(sp_pr, "a:custGeom")
        fill_kind, fill = self.fill_of(sp_pr, style)
        stroke, stroke_w, dash = self.line_of(sp_pr, style)
        if _local(sp) == "cxnSp" and prst is None:
            prst = "line"
        tx = _find(sp, "p:txBody")
        has_text = tx is not None and "".join(t.text or "" for t in tx.iter(_tag("a", "t"))).strip() != ""
        z = len(out) + 1

        drew_background = False
        if fill_kind == "blip":
            blip = _find(fill, "a:blip")
            src = self._blip_src(blip, fill, prst if prst in ("ellipse", "roundRect") else None)
            if src:
                out.append(dict(b, id=self.imp.eid("img"), type="image", src=src, alt="", zIndex=z,
                                opacity=self._blip_alpha(blip)))
                drew_background = True
        elif cust is not None or prst not in (None, "rect", "roundRect", "ellipse", "snipRoundRect") or fill_kind == "grad":
            # Keep single-colour PowerPoint freeforms and preset shapes editable.
            # They used to be flattened into an SVG image, which meant users could
            # not change their fill, outline or outline width in the editor.
            editable_paths = None
            if fill_kind in (None, "none", "solid"):
                editable_paths = self.editable_paths(prst, cust, b, x, geom)
            if editable_paths:
                fill_css = _css_color(fill) if fill_kind == "solid" and fill else "transparent"
                if fill_kind == "solid" and fill:
                    imp.fill_colors[fill[0]] += b["width"] * b["height"]
                out.append(dict(b, id=self.imp.eid("path"), type="shape", shapeType="path",
                                pathData=editable_paths, pathWidth=b["width"], pathHeight=b["height"],
                                pathFlipH=flip_h, pathFlipV=flip_v, fill=fill_css,
                                stroke=stroke or "transparent", strokeWidth=stroke_w if stroke else 0,
                                strokeStyle=dash, borderRadius=0, zIndex=z))
                drew_background = True
            else:
                svg_pad = self.shape_svg(prst, cust, fill_kind, fill, stroke, stroke_w, dash, b, flip_h, flip_v, x, geom,
                                         _find(sp_pr, "a:ln"))
                if svg_pad:
                    svg, pad = svg_pad
                    vb = dict(b, x=b["x"] - pad, y=b["y"] - pad, width=b["width"] + 2 * pad, height=b["height"] + 2 * pad)
                    out.append(dict(vb, id=self.imp.eid("vec"), type="image", src=imp.store_svg(svg), alt="", zIndex=z))
                    drew_background = True
                elif fill_kind == "solid" or stroke:
                    imp.warnings[f"shapes '{prst}' approximated as rectangles"] += 1
                    prst = "rect"
        if not drew_background and (fill_kind == "solid" and fill or stroke):
            fill_css = _css_color(fill) if fill_kind == "solid" and fill else "transparent"
            radius = 0
            if prst == "roundRect":
                adj = 16667
                gd = _find(geom, "a:avLst/a:gd")
                if gd is not None:
                    m = re.search(r"val\s+(-?\d+)", gd.get("fmla", ""))
                    adj = _int(m.group(1), 16667) if m else 16667
                radius = round(min(b["width"], b["height"]) * adj / 100000)
            if fill_kind == "solid" and fill:
                imp.fill_colors[fill[0]] += b["width"] * b["height"]
            out.append(dict(b, id=self.imp.eid("shp"), type="shape",
                            shapeType="circle" if prst == "ellipse" else "rectangle",
                            fill=fill_css, stroke=stroke or "transparent", strokeWidth=stroke_w if stroke else 0,
                            strokeStyle=dash, borderRadius=radius, zIndex=z))
        if has_text:
            text_el = self.text_element(sp, tx, ph, b)
            if text_el:
                text_el["zIndex"] = len(out) + 1
                text_el["id"] = self.imp.eid("txt")
                out.append(text_el)

    def editable_paths(self, prst, cust, b, xfrm_d, geom):
        """Return simple path geometry for an editable canvas vector shape.

        `pathData` intentionally contains only geometry and per-subpath fill/stroke
        flags; visual styling stays in the element model so editor controls can change
        it after import.
        """
        w, h = b["width"], b["height"]
        paths = []
        if cust is not None:
            for path in _findall(cust, "a:pathLst/a:path"):
                d = _cust_path_to_svg(path, w, h, xfrm_d["ext"])
                if d is None:
                    return None
                paths.append({"d": d, "noFill": path.get("fill") == "none",
                              "noStroke": path.get("stroke") in ("0", "false")})
        else:
            d = _preset_path(prst or "rect", w, h, geom)
            if d is None:
                return None
            paths.append({"d": d, "noFill": prst in ("line", "straightConnector1", "arc"), "noStroke": False})
        return paths or None

    def _blip_src(self, blip, blip_fill, mask: Optional[str] = None) -> Optional[str]:
        if blip is None:
            return None
        svg = blip.find(".//" + _tag("asvg", "svgBlip"))
        if svg is not None and svg.get(_tag("r", "embed")):
            url = self.imp.media_url(self.part, svg.get(_tag("r", "embed")))
            if url:
                return url
        r_id = blip.get(_tag("r", "embed"))
        if not r_id:
            if blip.get(_tag("r", "link")):
                self.imp.warnings["linked (external) pictures skipped"] += 1
            return None
        crop = None
        src_rect = _find(blip_fill, "a:srcRect")
        if src_rect is not None and src_rect.attrib:
            crop = tuple(_int(src_rect.get(k)) / 100000 for k in ("l", "t", "r", "b"))
        # Canva-style crop: <a:stretch><a:fillRect l t r b/> with negative values stretches the picture
        # beyond the frame and the frame cuts it. Convert it to the equivalent crop of the picture
        # (fraction of the stretched picture outside the frame), on top of any srcRect.
        fill_rect = _find(blip_fill, "a:stretch/a:fillRect")
        if fill_rect is not None and fill_rect.attrib:
            l, t, r, b = (_int(fill_rect.get(k)) / 100000 for k in ("l", "t", "r", "b"))
            span_x, span_y = 1 - l - r, 1 - t - b
            if span_x > 0 and span_y > 0 and min(l, t, r, b) < 0:
                extra = (max(0.0, -l) / span_x, max(0.0, -t) / span_y, max(0.0, -r) / span_x, max(0.0, -b) / span_y)
                base = crop or (0.0, 0.0, 0.0, 0.0)
                # compose: the srcRect part is applied first, the fillRect crop inside what remains
                w_left = 1 - base[0] - base[2]
                h_left = 1 - base[1] - base[3]
                crop = (base[0] + extra[0] * w_left, base[1] + extra[1] * h_left,
                        base[2] + extra[2] * w_left, base[3] + extra[3] * h_left)
        if crop and not any(abs(c) > 1e-4 for c in crop):
            crop = None
        return self.imp.media_url(self.part, r_id, crop, mask)

    @staticmethod
    def _blip_alpha(blip) -> float:
        amt = _find(blip, "a:alphaModFix")
        return round(_int(amt.get("amt"), 100000) / 100000, 3) if amt is not None else 1

    def picture(self, pic, ph, transform: Transform, out):
        imp = self.imp
        x = self.xfrm_for(pic, ph)
        if x is None:
            return
        b, _, _ = imp.box(x, transform)
        nv_pr = _find(pic, "p:nvPicPr/p:nvPr")
        media = None
        if nv_pr is not None:
            for node in nv_pr.iter(_tag("a", "videoFile"), _tag("p14", "media")):
                media = node.get(_tag("r", "embed")) or node.get(_tag("r", "link"))
                if media:
                    break
        if media:
            try:
                part = self.part.related_part(media)
                url = imp.store(part.blob, str(part.partname.ext).lower())
                out.append(dict(b, id=self.imp.eid("vid"), type="video", src=url, videoType="file", zIndex=len(out) + 1))
                return
            except Exception:
                imp.warnings["videos skipped (linked, unreadable or unsupported format)"] += 1
        blip_fill = _find(pic, "p:blipFill")
        blip = _find(blip_fill, "a:blip")
        geom = _find(pic, "p:spPr/a:prstGeom")
        mask = geom.get("prst") if geom is not None and geom.get("prst") in ("ellipse", "roundRect") else None
        src = self._blip_src(blip, blip_fill, mask)
        if src:
            out.append(dict(b, id=self.imp.eid("pic"), type="image", src=src, alt="", zIndex=len(out) + 1,
                            opacity=self._blip_alpha(blip)))

    def graphic_frame(self, gf, transform: Transform, out):
        imp = self.imp
        data = gf.find(".//" + _tag("a", "graphicData"))
        uri = data.get("uri", "") if data is not None else ""
        tbl = _find(data, "a:tbl")
        if tbl is None:
            kind = "charts" if "chart" in uri else ("SmartArt diagrams" if "diagram" in uri else "embedded objects")
            imp.warnings[f"{kind} skipped (not supported)"] += 1
            return
        x = _read_xfrm(_find(gf, "p:xfrm"))
        if x is None:
            return
        b, _, _ = imp.box(x, transform)
        widths = [_int(gc.get("w")) for gc in _findall(tbl, "a:tblGrid/a:gridCol")]
        rows = []
        for tr in _findall(tbl, "a:tr"):
            rows.append(["\n".join(_paragraph_text(p) for p in tc.iter(_tag("a", "p"))).strip()
                         for tc in _findall(tr, "a:tc")])
        if not rows:
            return
        ncols = max(len(r) for r in rows)
        rows = [r + [""] * (ncols - len(r)) for r in rows]
        trs = _findall(tbl, "a:tr")

        def cell_fill(tr_idx):
            if tr_idx >= len(trs):
                return None
            tc = _find(trs[tr_idx], "a:tc")
            c = self.colors().resolve(_find(tc, "a:tcPr/a:solidFill")) if tc is not None else None
            return _css_color(c) if c else None

        first_run = tbl.find(".//" + _tag("a", "rPr"))
        size_pt = _int(first_run.get("sz"), 1400) / 100 if first_run is not None else 14
        header_color = None
        if first_run is not None:
            c = self.colors().resolve(_find(first_run, "a:solidFill"))
            header_color = c[0] if c else None
        total = sum(widths) or 1
        out.append(dict(
            b, id=imp.eid("tbl"), type="table", columns=rows[0], rows=rows[1:] or [[""] * ncols],
            colWidths=[round(w / total * 100, 2) for w in widths] if len(widths) == ncols else None,
            headerFill=cell_fill(0) or "#1e293b", headerColor=header_color or "#ffffff",
            rowFill=cell_fill(1) or "transparent", altRowFill=cell_fill(2) or cell_fill(1) or "transparent",
            borderColor="#94a3b8", fontSize=max(10, round(size_pt * imp.px_per_pt)), zIndex=len(out) + 1))

    # ── text ──
    def level_styles(self, ph, lvl: int, tx) -> List[Any]:
        """Candidate <a:lvlNpPr> nodes, closest first: shape, layout/master placeholder, master text styles."""
        tag = f"a:lvl{lvl + 1}pPr"
        chain = [_find(tx, f"a:lstStyle/{tag}")]
        if ph is not None:
            for inh in self.inherited_ph(ph):
                chain.append(_find(inh, f"p:txBody/a:lstStyle/{tag}"))
        master_root = (self.master_ctx or self).root
        tx_styles = _find(master_root, "p:txStyles")
        p_type = ph.get("type", "body") if ph is not None else None
        style_name = "p:titleStyle" if p_type in ("title", "ctrTitle") else \
            ("p:bodyStyle" if ph is not None else "p:otherStyle")
        chain.append(_find(tx_styles, f"{style_name}/{tag}"))
        if ph is None:
            # Text boxes: the presentation-wide default text style applies
            pres = self.imp.prs.part._element
            chain.append(_find(pres, f"p:defaultTextStyle/{tag}"))
        return [c for c in chain if c is not None]

    def text_element(self, sp, tx, ph, b) -> Optional[Dict[str, Any]]:
        imp = self.imp
        paragraphs = _findall(tx, "a:p")
        body_pr = _find(tx, "a:bodyPr")
        font_scale = 1.0
        autofit = _find(body_pr, "a:normAutofit")
        if autofit is not None and autofit.get("fontScale"):
            font_scale = _int(autofit.get("fontScale"), 100000) / 100000
        style = _find(sp, "p:style")
        font_ref_color = self.colors().resolve(_find(style, "a:fontRef")) if style is not None else None

        is_title = ph is not None and ph.get("type") in ("title", "ctrTitle")
        texts, bullet_flags, numbered = [], [], False
        run_stats: Counter = Counter()
        align, line_pct = None, None
        for p in paragraphs:
            text = _paragraph_text(p)
            texts.append(text)
            p_pr = _find(p, "a:pPr")
            lvl = _int(p_pr.get("lvl")) if p_pr is not None else 0
            chain = ([p_pr] if p_pr is not None else []) + self.level_styles(ph, lvl, tx)
            if text.strip():
                bullet = None
                for node in chain:
                    if _find(node, "a:buNone") is not None:
                        bullet = False
                        break
                    if _find(node, "a:buAutoNum") is not None:
                        bullet, numbered = True, True
                        break
                    if _find(node, "a:buChar") is not None or _find(node, "a:buBlip") is not None:
                        bullet = True
                        break
                bullet_flags.append(bool(bullet))
                if align is None:
                    for node in chain:
                        if node.get("algn"):
                            align = node.get("algn")
                            break
                if line_pct is None:
                    for node in chain:
                        sp_pct = _find(node, "a:lnSpc/a:spcPct")
                        if sp_pct is not None:
                            line_pct = _int(sp_pct.get("val"), 100000) / 100000
                            break
            def_rprs = [_find(n, "a:defRPr") for n in chain]
            def_rprs = [d for d in def_rprs if d is not None]
            for r in p.iter(_tag("a", "r")):
                t = "".join(x.text or "" for x in r.iter(_tag("a", "t")))
                if not t.strip():
                    continue
                r_pr = _find(r, "a:rPr")
                props = self._run_props(r_pr, def_rprs, font_ref_color, is_title)
                run_stats[props] += len(t.strip())

        if not any(t.strip() for t in texts):
            return None
        props = run_stats.most_common(1)[0][0] if run_stats else ("", 18.0, False, False, None, False)
        font_name, size_pt, bold, italic, color, underline = props
        clean, weight, name_italic = _clean_font_name(font_name)
        size_px = max(8, round(size_pt * font_scale * imp.px_per_pt))
        for (fn, sz, *_rest), n in run_stats.items():
            imp.run_fonts[_clean_font_name(fn)[0]] += n
            if sz >= 28:
                imp.big_run_fonts[_clean_font_name(fn)[0]] += n
        if color:
            imp.text_colors[color[0]] += sum(run_stats.values())
        css_color = _css_color(color) if color else None
        gradient = self._text_gradient(tx)
        if gradient:
            css_color = gradient[0]

        element = dict(b)
        common = {
            "fontFamily": _resolve_font(clean)[0] if clean else None,
            "fontSize": size_px,
            "color": css_color,
            "lineHeight": round(1.2 * line_pct, 2) if line_pct else 1.2,
            "colorGradient": gradient,
        }
        common = {k: v for k, v in common.items() if v is not None}
        non_empty = [t for t in texts if t.strip()]
        if len(non_empty) >= 2 and all(bullet_flags):
            element.update(type="bullets", items=non_empty, bulletStyle="number" if numbered else "disc",
                           autoHeight=False, **common)
            self._fit_insets(element, body_pr)
            return element
        element.update(
            type="text",
            content="\n".join(texts).strip("\n"),
            fontWeight=max(weight or 400, 700 if bold else 400),
            fontStyle="italic" if italic or name_italic else "normal",
            textDecoration="underline" if underline else "none",
            textAlign={"ctr": "center", "r": "right", "just": "justify", "dist": "justify"}.get(align or "l", "left"),
            **common,
        )
        self._apply_vertical_anchor(element, body_pr, ph, len(texts))
        self._fit_insets(element, body_pr)
        return element

    def _run_props(self, r_pr, def_rprs, font_ref_color, is_title):
        nodes = ([r_pr] if r_pr is not None else []) + def_rprs

        def attr(name):
            for n in nodes:
                if n.get(name) is not None:
                    return n.get(name)
            return None

        size = _int(attr("sz"), 1800) / 100
        bold = attr("b") in ("1", "true")
        italic = attr("i") in ("1", "true")
        underline = attr("u") not in (None, "none")
        font = None
        for n in nodes:
            latin = _find(n, "a:latin")
            if latin is not None and latin.get("typeface"):
                font = latin.get("typeface")
                break
        if font is None:
            font = self.imp.major_font if is_title else self.imp.minor_font
        elif font == "+mn-lt":
            font = self.imp.minor_font
        elif font == "+mj-lt":
            font = self.imp.major_font
        color = None
        for n in nodes:
            solid = _find(n, "a:solidFill")
            if solid is not None:
                color = self.colors().resolve(solid)
                break
        if color is None:
            color = font_ref_color
        return (font or "", size, bold, italic, color, underline)

    def _text_gradient(self, body):
        """Gradient text fill (Canva numbers): the distinct stop colours of the first gradient run, or None."""
        if body is None:
            return None
        for r_pr in body.iter(_tag("a", "rPr")):
            grad = _find(r_pr, "a:gradFill")
            if grad is None:
                continue
            stops = []
            for gs in _findall(grad, "a:gsLst/a:gs"):
                c = self.colors().resolve(gs)
                if c and _css_color(c) not in stops:
                    stops.append(_css_color(c))
            if len(stops) >= 2:
                return stops
        return None

    EDITOR_PAD_X, EDITOR_PAD_Y = 12, 8

    def _fit_insets(self, element, body_pr):
        """Move and grow the box so that the editor's fixed text padding lands where the PPTX
        text insets were, and leave room for wider substitute fonts."""
        s = self.imp.scale
        l_ins = _int(body_pr.get("lIns"), 91440) * s if body_pr is not None else 91440 * s
        r_ins = _int(body_pr.get("rIns"), 91440) * s if body_pr is not None else 91440 * s
        t_ins = _int(body_pr.get("tIns"), 45720) * s if body_pr is not None else 45720 * s
        grow_l = self.EDITOR_PAD_X - l_ins
        grow_r = self.EDITOR_PAD_X - r_ins
        element["x"] = round(element["x"] - grow_l)
        element["width"] = round(element["width"] + grow_l + grow_r)
        element["y"] = round(element["y"] - (self.EDITOR_PAD_Y - t_ins))
        element["height"] = round(element["height"] + max(0, 2 * self.EDITOR_PAD_Y - t_ins))
        # Fonts missing on this machine are replaced by wider ones: allow 12% (30% if the box never wraps),
        # and enough width for the longest word, which fitted on one line in PowerPoint
        no_wrap = body_pr is not None and body_pr.get("wrap") == "none"
        extra = round(element["width"] * (0.30 if no_wrap else 0.12))
        text = element.get("content") or " ".join(element.get("items", []))
        longest = max((len(w) for w in text.split()), default=0)
        char_w = 0.62 if (element.get("fontWeight") or 400) >= 600 else 0.56
        needed = round(longest * element["fontSize"] * char_w) + 2 * self.EDITOR_PAD_X
        extra = max(extra, needed - element["width"])
        align = element.get("textAlign", "left")
        if align == "center":
            element["x"] -= extra // 2
        elif align == "right":
            element["x"] -= extra
        element["width"] += extra

    def _apply_vertical_anchor(self, element, body_pr, ph, n_paragraphs):
        """The editor draws text from the top; emulate centered or bottom anchoring by
        shrinking the box around the estimated text height."""
        anchor = body_pr.get("anchor") if body_pr is not None else None
        if anchor is None and ph is not None:
            for inh in self.inherited_ph(ph):
                bp = _find(inh, "p:txBody/a:bodyPr")
                if bp is not None and bp.get("anchor"):
                    anchor = bp.get("anchor")
                    break
        if anchor not in ("ctr", "b") or element.get("rotation"):
            return
        size = element["fontSize"]
        width = max(40, element["width"] - 24)
        lines = 0
        for para in element["content"].split("\n"):
            lines += max(1, math.ceil(len(para) * size * 0.55 / width))
        est = round(lines * size * element["lineHeight"] + 20)
        if est >= element["height"]:
            return
        if anchor == "ctr":
            element["y"] += round((element["height"] - est) / 2)
        else:
            element["y"] += element["height"] - est
        element["height"] = est

    # ── vector shapes as SVG ──
    def shape_svg(self, prst, cust, fill_kind, fill, stroke, stroke_w, dash, b, flip_h, flip_v, xfrm_d, geom, ln):
        """SVG for shapes the editor cannot draw natively. Returns (svg, pad): the SVG is `pad`
        pixels larger on each side than the shape so that strokes and arrowheads are not clipped."""
        w, h = b["width"], b["height"]
        if w < 1 or h < 1:
            return None
        defs: List[str] = []
        fill_attr = 'fill="none"'
        if fill_kind == "solid" and fill:
            fill_attr = f'fill="{fill[0]}" fill-opacity="{round(fill[1], 3)}"'
        elif fill_kind == "grad":
            stops = []
            for gs in _findall(fill, "a:gsLst/a:gs"):
                c = self.colors().resolve(gs)
                if c:
                    stops.append(f'<stop offset="{_int(gs.get("pos")) / 1000:.1f}%" stop-color="{c[0]}" '
                                 f'stop-opacity="{round(c[1], 3)}"/>')
            if not stops:
                return None
            lin = _find(fill, "a:lin")
            rad = math.radians(_int(lin.get("ang")) / 60000 if lin is not None else 90)
            x2, y2 = 50 + 50 * math.cos(rad), 50 + 50 * math.sin(rad)
            defs.append(f'<linearGradient id="g" x1="{100 - x2:.1f}%" y1="{100 - y2:.1f}%" x2="{x2:.1f}%" '
                        f'y2="{y2:.1f}%">' + "".join(stops) + "</linearGradient>")
            fill_attr = 'fill="url(#g)"'
        sw = stroke_w if stroke else 0
        if fill_attr == 'fill="none"' and not stroke:
            return None

        stroke_attrs = f'stroke="{stroke}" stroke-width="{sw}" stroke-linejoin="round"' if stroke else 'stroke="none"'
        if stroke and dash == "dashed":
            stroke_attrs += f' stroke-dasharray="{sw * 4},{sw * 3}"'
        elif stroke and dash == "dotted":
            stroke_attrs += f' stroke-dasharray="{sw},{sw * 2}"'
        markers = ""
        for end, attr in (("headEnd", "marker-start"), ("tailEnd", "marker-end")):
            node = _find(ln, f"a:{end}")
            if stroke and node is not None and node.get("type", "none") != "none":
                orient = "auto-start-reverse" if end == "headEnd" else "auto"
                defs.append(f'<marker id="m_{end}" viewBox="0 0 10 10" refX="8" refY="5" markerWidth="4" '
                            f'markerHeight="4" orient="{orient}"><path d="M0 0 L10 5 L0 10 Z" fill="{stroke}"/></marker>')
                markers += f' {attr}="url(#m_{end})"'

        paths = []
        if cust is not None:
            for path in _findall(cust, "a:pathLst/a:path"):
                d = _cust_path_to_svg(path, w, h, xfrm_d["ext"])
                if d is None:
                    self.imp.warnings["freeform shapes skipped (formula-based geometry)"] += 1
                    return None
                paths.append((d, path.get("fill") == "none", path.get("stroke") in ("0", "false")))
        else:
            d = _preset_path(prst or "rect", w, h, geom)
            if d is None:
                return None
            paths.append((d, prst in ("line", "straightConnector1", "arc"), False))
        if not paths:
            return None

        pad = (math.ceil(sw * 2) if markers else math.ceil(sw / 2)) + 1 if sw else 0
        tx, ty = pad + (w if flip_h else 0), pad + (h if flip_v else 0)
        transform = f' transform="translate({tx},{ty})' + (
            f' scale({-1 if flip_h else 1},{-1 if flip_v else 1})"' if flip_h or flip_v else '"')
        no_fill_attr, no_stroke_attr = 'fill="none"', 'stroke="none"'
        body = "".join(
            f'<path d="{d}" {no_fill_attr if no_fill else fill_attr} '
            f'{no_stroke_attr if no_stroke else stroke_attrs + markers}/>'
            for d, no_fill, no_stroke in paths)
        W, H = w + 2 * pad, h + 2 * pad
        defs_xml = f"<defs>{''.join(defs)}</defs>" if defs else ""
        return (f'<svg xmlns="http://www.w3.org/2000/svg" width="{W}" height="{H}" viewBox="0 0 {W} {H}" '
                f'preserveAspectRatio="none">{defs_xml}<g{transform}>{body}</g></svg>'), pad


def _paragraph_text(p) -> str:
    parts = []
    for node in p:
        n = _local(node)
        if n in ("r", "fld"):
            parts.append("".join(t.text or "" for t in node.iter(_tag("a", "t"))))
        elif n == "br":
            parts.append("\n")
    return "".join(parts)


def _cust_path_to_svg(path, w, h, ext) -> Optional[str]:
    pw = _int(path.get("w")) or ext[0] or 1
    ph = _int(path.get("h")) or ext[1] or 1
    sx, sy = w / pw, h / ph
    cmds = []
    cur = (0.0, 0.0)

    def pt(node):
        x, y = node.get("x"), node.get("y")
        if x is None or y is None or not re.fullmatch(r"-?\d+", x) or not re.fullmatch(r"-?\d+", y):
            raise ValueError("formula")
        return int(x) * sx, int(y) * sy

    try:
        for node in path:
            n = _local(node)
            pts = [pt(p) for p in _findall(node, "a:pt")]
            if n == "moveTo":
                cur = pts[0]
                cmds.append(f"M{cur[0]:.2f} {cur[1]:.2f}")
            elif n == "lnTo":
                cur = pts[0]
                cmds.append(f"L{cur[0]:.2f} {cur[1]:.2f}")
            elif n == "cubicBezTo":
                cmds.append("C" + " ".join(f"{x:.2f} {y:.2f}" for x, y in pts))
                cur = pts[-1]
            elif n == "quadBezTo":
                cmds.append("Q" + " ".join(f"{x:.2f} {y:.2f}" for x, y in pts))
                cur = pts[-1]
            elif n == "arcTo":
                vals = [node.get(k) for k in ("wR", "hR", "stAng", "swAng")]
                if not all(v is not None and re.fullmatch(r"-?\d+", v) for v in vals):
                    raise ValueError("formula")
                wr, hr = int(vals[0]) * sx, int(vals[1]) * sy
                st, sw = math.radians(int(vals[2]) / 60000), math.radians(int(vals[3]) / 60000)
                cx, cy = cur[0] - wr * math.cos(st), cur[1] - hr * math.sin(st)
                end = (cx + wr * math.cos(st + sw), cy + hr * math.sin(st + sw))
                large = 1 if abs(sw) > math.pi else 0
                sweep = 1 if sw > 0 else 0
                cmds.append(f"A{wr:.2f} {hr:.2f} 0 {large} {sweep} {end[0]:.2f} {end[1]:.2f}")
                cur = end
            elif n == "close":
                cmds.append("Z")
    except (ValueError, IndexError):
        return None
    return " ".join(cmds) if cmds else None


def _preset_path(prst: str, w: float, h: float, geom) -> Optional[str]:
    """SVG paths for the common preset geometries; others are approximated by the caller."""
    if prst in ("line", "straightConnector1"):
        return f"M0 0 L{w} {h}"
    if prst == "rect":
        return f"M0 0 H{w} V{h} H0 Z"
    if prst == "roundRect":
        r = min(w, h) * 0.16667
        return (f"M{r} 0 H{w - r} A{r} {r} 0 0 1 {w} {r} V{h - r} A{r} {r} 0 0 1 {w - r} {h} "
                f"H{r} A{r} {r} 0 0 1 0 {h - r} V{r} A{r} {r} 0 0 1 {r} 0 Z")
    if prst == "ellipse":
        return f"M0 {h / 2} A{w / 2} {h / 2} 0 1 1 {w} {h / 2} A{w / 2} {h / 2} 0 1 1 0 {h / 2} Z"
    if prst == "triangle":
        return f"M{w / 2} 0 L{w} {h} L0 {h} Z"
    if prst == "rtTriangle":
        return f"M0 0 L{w} {h} L0 {h} Z"
    if prst == "diamond":
        return f"M{w / 2} 0 L{w} {h / 2} L{w / 2} {h} L0 {h / 2} Z"
    if prst == "parallelogram":
        o = min(w, h) * 0.25
        return f"M{o} 0 L{w} 0 L{w - o} {h} L0 {h} Z"
    if prst == "trapezoid":
        o = min(w, h) * 0.25
        return f"M{o} 0 L{w - o} 0 L{w} {h} L0 {h} Z"
    if prst == "hexagon":
        o = min(w, h) * 0.25
        return f"M{o} 0 L{w - o} 0 L{w} {h / 2} L{w - o} {h} L{o} {h} L0 {h / 2} Z"
    if prst == "octagon":
        o = min(w, h) * 0.29
        return f"M{o} 0 L{w - o} 0 L{w} {o} L{w} {h - o} L{w - o} {h} L{o} {h} L0 {h - o} L0 {o} Z"
    if prst == "pentagon":
        return f"M{w / 2} 0 L{w} {h * 0.38} L{w * 0.81} {h} L{w * 0.19} {h} L0 {h * 0.38} Z"
    if prst == "homePlate":
        o = min(w, h) * 0.5
        return f"M0 0 L{w - o} 0 L{w} {h / 2} L{w - o} {h} L0 {h} Z"
    if prst == "chevron":
        o = min(w, h) * 0.5
        return f"M0 0 L{w - o} 0 L{w} {h / 2} L{w - o} {h} L0 {h} L{o} {h / 2} Z"
    if prst == "rightArrow":
        hw, sh = min(w, h) * 0.5, h * 0.25
        return f"M0 {sh} L{w - hw} {sh} L{w - hw} 0 L{w} {h / 2} L{w - hw} {h} L{w - hw} {h - sh} L0 {h - sh} Z"
    if prst == "leftArrow":
        hw, sh = min(w, h) * 0.5, h * 0.25
        return f"M{w} {sh} L{hw} {sh} L{hw} 0 L0 {h / 2} L{hw} {h} L{hw} {h - sh} L{w} {h - sh} Z"
    if prst == "plus":
        o = min(w, h) * 0.25
        return (f"M{o} 0 H{w - o} V{o} H{w} V{h - o} H{w - o} V{h} H{o} V{h - o} H0 V{o} H{o} Z")
    if prst in ("snipRoundRect", "round2SameRect", "round1Rect", "snip1Rect", "snip2SameRect", "flowChartProcess",
                "flowChartAlternateProcess", "rect"):
        return f"M0 0 H{w} V{h} H0 Z"
    return None


# ── theme, backgrounds, palette ────────────────────────────────────────
def _theme_from_master(prs) -> Optional[Any]:
    try:
        from pptx.opc.constants import RELATIONSHIP_TYPE as RT
        master = prs.slide_masters[0]
        return etree.fromstring(master.part.part_related_by(RT.THEME).blob)
    except Exception:
        return None


def _clr_map(prs) -> Dict[str, str]:
    try:
        cm = _find(prs.slide_masters[0]._element, "p:clrMap")
        if cm is not None:
            return {k: v for k, v in cm.attrib.items()}
    except Exception:
        pass
    return {}


def _background_of(ctx: "_Ctx", theme_tree) -> Optional[Dict[str, Any]]:
    """Resolve <p:bg> of one part: bgPr fill, or bgRef into the theme's background fill styles."""
    bg = ctx.root.find(".//" + _tag("p", "bg"))
    if bg is None:
        return None
    resolver = ctx.colors()
    bg_pr = _find(bg, "p:bgPr")
    fill_node = None
    ph_color = None
    if bg_pr is not None:
        fill_node = bg_pr
    else:
        ref = _find(bg, "p:bgRef")
        if ref is not None and theme_tree is not None:
            ph_color = resolver.resolve(ref)
            idx = _int(ref.get("idx"))
            lst = theme_tree.find(".//" + _tag("a", "bgFillStyleLst")) if idx >= 1001 else \
                theme_tree.find(".//" + _tag("a", "fillStyleLst"))
            if lst is not None:
                items = list(lst)
                pos = idx - 1001 if idx >= 1001 else idx - 1
                if 0 <= pos < len(items):
                    fill_node = etree.Element("wrap")
                    fill_node.append(_copy(items[pos]))
    if fill_node is None:
        return None
    for node in fill_node:
        n = _local(node)
        if n == "solidFill":
            c = resolver.resolve(node, ph_color)
            if c:
                return {"type": "solid", "color": c[0]}
        if n == "gradFill":
            stops = []
            for gs in _findall(node, "a:gsLst/a:gs"):
                c = resolver.resolve(gs, ph_color)
                if c:
                    stops.append({"color": c[0], "position": _int(gs.get("pos")) / 1000})
            if stops:
                return {"type": "gradient", "stops": stops}
        if n == "blipFill":
            blip = _find(node, "a:blip")
            r_id = blip.get(_tag("r", "embed")) if blip is not None else None
            if r_id:
                try:
                    part = ctx.part.related_part(r_id)
                    ext = str(part.partname.ext).lower()
                    mime = {"png": "image/png", "jpg": "image/jpeg", "jpeg": "image/jpeg", "gif": "image/gif",
                            "svg": "image/svg+xml", "webp": "image/webp"}.get(ext)
                    if mime:
                        return {"type": "image", "dataurl": f"data:{mime};base64,{base64.b64encode(part.blob).decode()}",
                                "_key": str(part.partname)}
                except Exception:
                    pass
    return None


def _copy(el):
    return etree.fromstring(etree.tostring(el))


def _bg_signature(bg: Optional[Dict[str, Any]]) -> Optional[str]:
    if not bg:
        return None
    if bg["type"] == "image":
        return "image:" + bg.get("_key", bg.get("dataurl", "")[:200])
    if bg["type"] == "gradient":
        return "grad:" + repr(bg["stops"])
    return "solid:" + bg.get("color", "")


def _average_color(dataurl: str) -> Optional[str]:
    try:
        from PIL import Image
        raw = base64.b64decode(dataurl.split(",", 1)[1])
        img = Image.open(io.BytesIO(raw)).convert("RGB").resize((1, 1))
        return _hex(img.getpixel((0, 0)))
    except Exception:
        return None


def _bg_element(bg: Dict[str, Any], imp: _Importer) -> Optional[Dict[str, Any]]:
    """A full-slide element showing a slide-specific background."""
    full = {"x": 0, "y": 0, "width": CANVAS_W, "height": CANVAS_H, "rotation": 0, "zIndex": 0}
    if bg["type"] == "solid":
        return dict(full, id="el_bg_solid", type="shape", shapeType="rectangle", fill=bg["color"],
                    stroke="transparent", strokeWidth=0, strokeStyle="solid", borderRadius=0)
    if bg["type"] == "image":
        mime, b64 = bg["dataurl"][5:].split(";base64,", 1)
        ext = {"image/jpeg": "jpg", "image/svg+xml": "svg"}.get(mime, mime.split("/")[-1])
        try:
            src = imp.store(base64.b64decode(b64), ext)
        except ValueError as e:
            imp.warnings[f"slide backgrounds skipped ({e})"] += 1
            return None
        return dict(full, id="el_bg_image", type="image", src=src, alt="")
    stops = "".join(f'<stop offset="{s["position"]:.1f}%" stop-color="{s["color"]}"/>' for s in bg["stops"])
    svg = (f'<svg xmlns="http://www.w3.org/2000/svg" width="{CANVAS_W}" height="{CANVAS_H}" preserveAspectRatio="none">'
           f'<defs><linearGradient id="g" x1="0" y1="0" x2="1" y2="1">{stops}</linearGradient></defs>'
           f'<rect width="100%" height="100%" fill="url(#g)"/></svg>')
    return dict(full, id="el_bg_grad", type="image", src=imp.store_svg(svg), alt="")


def _data_url_store(blob: bytes, ext: str) -> str:
    mime = {"png": "image/png", "jpg": "image/jpeg", "jpeg": "image/jpeg", "gif": "image/gif",
            "svg": "image/svg+xml", "webp": "image/webp", "mp4": "video/mp4", "webm": "video/webm",
            "mov": "video/quicktime", "m4v": "video/mp4"}.get(ext, "application/octet-stream")
    return f"data:{mime};base64,{base64.b64encode(blob).decode('utf-8')}"


# ── main entry point ───────────────────────────────────────────────────
def extract_theme(pptx_bytes: bytes, store_media: Optional[Callable[[bytes, str], str]] = None) -> dict:
    """Extract theme colors, fonts, background and slides from PPTX bytes (see module docstring)."""
    store_media = store_media or _data_url_store
    result: Dict[str, Any] = {
        "colors": {},
        "fonts": {"heading": {}, "body": {}},
        "background": {"type": "solid", "color": "#0f172a"},
        "warnings": [],
    }
    try:
        zipfile.ZipFile(io.BytesIO(pptx_bytes))
    except zipfile.BadZipFile:
        return result
    try:
        from pptx import Presentation
        prs = Presentation(io.BytesIO(pptx_bytes))
    except Exception as e:
        result["warnings"].append(f"Could not open the presentation: {e}")
        return result

    # 1. Theme: color scheme and fonts of the slide master's theme
    theme_tree = _theme_from_master(prs)
    scheme: Dict[str, str] = {}
    major_font = minor_font = ""
    if theme_tree is not None:
        clr_scheme = theme_tree.find(".//" + _tag("a", "clrScheme"))
        base = ColorResolver({})
        for key in _SCHEME_ORDER:
            el = _find(clr_scheme, f"a:{key}")
            c = base.resolve(el) if el is not None else None
            if c:
                scheme[key] = c[0]
                result["colors"][_SCHEME_LABELS[key]] = c[0]
        font_scheme = theme_tree.find(".//" + _tag("a", "fontScheme"))
        major = _find(font_scheme, "a:majorFont/a:latin")
        minor = _find(font_scheme, "a:minorFont/a:latin")
        major_font = major.get("typeface", "") if major is not None else ""
        minor_font = minor.get("typeface", "") if minor is not None else ""

    imp = _Importer(prs, store_media, scheme, major_font, minor_font)
    imp.color_resolver = ColorResolver(scheme, _clr_map(prs) or None)

    # 2. Backgrounds: per slide (slide, then layout, then master); the most common one becomes the theme
    slide_bgs = []
    for slide in prs.slides:
        layout = slide.slide_layout
        master = layout.slide_master
        ctx_m = _Ctx(imp, master.part, master._element, None, None, "master")
        ctx_l = _Ctx(imp, layout.part, layout._element, None, ctx_m, "layout")
        ctx_s = _Ctx(imp, slide.part, slide._element, ctx_l, ctx_m, "slide")
        slide_bgs.append(_background_of(ctx_s, theme_tree) or _background_of(ctx_l, theme_tree)
                         or _background_of(ctx_m, theme_tree))
    if not slide_bgs and len(prs.slide_masters):
        m = prs.slide_masters[0]
        slide_bgs.append(_background_of(_Ctx(imp, m.part, m._element, None, None, "master"), theme_tree))

    # 3. Slides
    extracted_slides = []
    n_slides = len(prs.slides)
    for s_idx, slide in enumerate(prs.slides):
        try:
            elements = imp.import_slide(slide)
        except Exception:
            imp.warnings["slides could not be imported"] += 1
            elements = []
        extracted_slides.append(elements)

    # Theme background: the most common slide background. Without any <p:bg> (typical of Canva
    # exports), the most common full-slide bottom element becomes the background instead.
    sigs = Counter(_bg_signature(b) for b in slide_bgs if b)
    theme_bg = None
    if sigs:
        top = sigs.most_common(1)[0][0]
        theme_bg = next(b for b in slide_bgs if b and _bg_signature(b) == top)
    else:
        bottoms = []
        for els in extracted_slides:
            first = els[0] if els else None
            if first and imp.is_full_bleed(first) and (first["type"] == "image" or
                                                       (first["type"] == "shape" and first.get("fill", "").startswith("#"))):
                bottoms.append(first.get("src") or first.get("fill"))
        if bottoms:
            common, count = Counter(bottoms).most_common(1)[0]
            if count >= max(1, len(extracted_slides) // 2):
                if common.startswith("#"):
                    theme_bg = {"type": "solid", "color": common}
                else:
                    theme_bg = {"type": "image", "src": common}
                for els in extracted_slides:
                    if els and imp.is_full_bleed(els[0]) and (els[0].get("src") or els[0].get("fill")) == common:
                        els.pop(0)
    if theme_bg:
        if theme_bg["type"] == "image" and "src" in theme_bg:
            # One image only: kept as a data URL because the PPTX exporter embeds it from there
            result["background"] = {"type": "image", "dataurl": imp.as_data_url(theme_bg["src"])}
        else:
            result["background"] = {k: v for k, v in theme_bg.items() if not k.startswith("_")}
    # Slides whose background differs from the theme background get it as a bottom element
    theme_sig = _bg_signature(theme_bg) if theme_bg and "src" not in theme_bg else None
    for els, bg in zip(extracted_slides, slide_bgs):
        if bg and _bg_signature(bg) != theme_sig:
            el = _bg_element(bg, imp)
            if el:
                els.insert(0, el)

    slides_out = []
    for s_idx, els in enumerate(extracted_slides):
        for z, el in enumerate(els, start=1):
            el["zIndex"] = z
            if el.get("opacity") == 1:
                el.pop("opacity")
        if s_idx == 0:
            slide_type, label = "title", "Diapositive de titre"
        elif s_idx == n_slides - 1:
            slide_type, label = "thankyou", "Merci / Fin"
        else:
            slide_type, label = "blank", f"Diapositive {s_idx + 1}"
        slides_out.append({"type": slide_type, "label": label, "elements": els})
    if slides_out:
        result["slides"] = slides_out

    # 4. Fonts: what the slides really use; the theme fonts when slides have no text
    def font_entry(name):
        clean, weight, _ = _clean_font_name(name)
        stack, default_weight = _resolve_font(clean)
        return {"name": clean, "stack": stack, "weight": weight or default_weight}

    body_font = imp.run_fonts.most_common(1)[0][0] if imp.run_fonts else minor_font
    heading_font = imp.big_run_fonts.most_common(1)[0][0] if imp.big_run_fonts else (major_font or body_font)
    if body_font:
        result["fonts"]["body"] = font_entry(body_font)
    if heading_font:
        result["fonts"]["heading"] = font_entry(heading_font)

    # 5. Palette
    bg_info = result["background"]
    if bg_info["type"] == "solid":
        bg_color = bg_info["color"]
    elif bg_info["type"] == "gradient":
        bg_color = bg_info["stops"][0]["color"]
    elif bg_info.get("dataurl"):
        bg_color = _average_color(bg_info["dataurl"]) or scheme.get("lt1", "#ffffff")
    else:
        bg_color = scheme.get("lt1", "#ffffff")
    if imp.text_colors:
        text_color = imp.text_colors.most_common(1)[0][0]
        if _contrast(text_color, bg_color) < 3:
            text_color = None
    else:
        text_color = None
    if not text_color:
        candidates = [scheme.get("dk1", "#000000"), scheme.get("lt1", "#ffffff"), "#0f172a", "#ffffff"]
        text_color = max(candidates, key=lambda c: _contrast(c, bg_color))
    accents = [scheme.get(f"accent{i}") for i in range(1, 7) if scheme.get(f"accent{i}")]
    if not accents or accents[0] in _OFFICE_DEFAULT_ACCENT1:
        used = [c for c, _ in (imp.fill_colors + imp.text_colors).most_common(30)
                if not _is_neutral(c) and _contrast(c, bg_color) > 1.3]
        distinct: List[str] = []
        for c in used:
            if all(sum(abs(a - b) for a, b in zip(_rgb(c), _rgb(d))) > 60 for d in distinct):
                distinct.append(c)
        accents = distinct + [a for a in accents if a not in distinct]
    accents += ["#3b82f6", "#10b981", "#8b5cf6", "#f43f5e", "#06b6d4"]
    r1, g1, b1 = _rgb(text_color)
    r2, g2, b2 = _rgb(bg_color)
    result["palette"] = {
        "bg": bg_color,
        "surface": scheme.get("lt1", "#ffffff"),
        "text": text_color,
        "text_muted": _hex(((r1 * 0.6 + r2 * 0.4), (g1 * 0.6 + g2 * 0.4), (b1 * 0.6 + b2 * 0.4))),
        "primary": accents[0],
        "secondary": accents[1],
        "accent": accents[2],
        "accent_2": accents[3],
        "accent_3": accents[4],
    }
    result["warnings"] = [f"{n} {what}" for what, n in imp.warnings.most_common()]
    return result
